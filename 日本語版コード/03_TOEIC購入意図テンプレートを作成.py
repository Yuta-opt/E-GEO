from __future__ import annotations

import argparse
import json
import random
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_OUTPUT_DIR = Path("日本語版データ/TOEIC/03_購入意図")
DEFAULT_ARCHIVE_DIR = Path("日本語版データ/TOEIC/99_旧版/03_購入意図")
DEFAULT_SEED = 20260731

CATEGORY_QUOTAS = {
    "general": 16,
    "vocabulary": 16,
    "grammar": 14,
    "mock_test": 14,
    "listening": 10,
    "reading": 10,
}
SPLIT_CATEGORY_QUOTAS = {
    "train": {
        "general": 8, "vocabulary": 8, "grammar": 7,
        "mock_test": 7, "listening": 5, "reading": 5,
    },
    "validation": {
        "general": 2, "vocabulary": 2, "grammar": 2,
        "mock_test": 2, "listening": 1, "reading": 1,
    },
    "test": {
        "general": 6, "vocabulary": 6, "grammar": 5,
        "mock_test": 5, "listening": 4, "reading": 4,
    },
}

SCORE_PATHS = [
    ("初受験", "500"), ("350点前後", "500"), ("400点前後", "600"),
    ("450点前後", "600"), ("500点前後", "650"), ("550点前後", "700"),
    ("600点前後", "730"), ("600点前後", "800"), ("650点前後", "800"),
    ("700点前後", "850"), ("730点前後", "860"), ("750点前後", "900"),
    ("800点前後", "900"), ("850点前後", "950"),
    ("社会人・500点前後", "700"), ("大学生・600点前後", "730"),
]
STUDY_PERIODS = ["1か月", "2か月", "3か月", "4か月", "6か月"]

CATEGORY_CONFIG = {
    "general": {
        "label": "総合対策",
        "short": "総合対策",
        "focus": "現在スコア、目標スコア、学習期間、全パート、一冊完結",
        "needs": [
            "全パートを基礎から学ぶ", "全パートをバランスよく伸ばす",
            "英語の基礎から立て直す", "苦手パートを見つけて補う",
            "600点台の壁を越える", "必要点を短期間で取る",
            "高得点帯の取りこぼしを減らす", "学業や仕事と両立する",
        ],
        "must": [
            "一冊で全体像を学べる", "初心者向けの丁寧な解説",
            "例題と演習の両方がある", "学習計画を立てやすい",
            "全パートを効率よく復習できる", "中上級者向けの総合演習",
        ],
        "nice": ["音声付き", "復習しやすい構成", "学習順が明確", "短時間で進めやすい"],
        "avoid": ["一分野だけに偏る教材", "解説が少なすぎる教材", "学習量が極端に多い教材"],
    },
    "vocabulary": {
        "label": "単語・熟語",
        "short": "単語",
        "focus": "目標スコア、語彙レベル、例文、音声、反復のしやすさ",
        "needs": [
            "基礎単語を覚える", "頻出単語を効率よく覚える",
            "単語とフレーズを一緒に覚える", "多義語と語法を学ぶ",
            "Part 7の言い換えに強くなる", "ビジネス語彙を増やす",
            "高得点向けの熟語を増やす", "毎日反復して定着させる",
        ],
        "must": [
            "例文が分かりやすい", "音声で反復できる",
            "目標点別に整理されている", "品詞や用法が分かる",
            "同義語や派生語も学べる", "復習チェックがある",
        ],
        "nice": ["持ち運びやすい", "アプリ対応", "確認問題付き", "短い単位で学べる"],
        "avoid": ["難語に偏る教材", "例文がほとんどない教材", "収録語が多すぎる教材"],
    },
    "grammar": {
        "label": "文法・Part 5/6",
        "short": "文法",
        "focus": "Part 5/6、苦手分野、問題量、解説、復習のしやすさ",
        "needs": [
            "英文法の基礎を学び直す", "Part 5の品詞問題を強化する",
            "時制・態・一致のミスを減らす", "Part 5と6をまとめて伸ばす",
            "語法と前置詞の弱点を補う", "Part 6の文脈問題を強化する",
            "Part 5を短時間で解く", "高難度の文法問題に対応する",
        ],
        "must": [
            "誤答理由まで解説する", "項目別に整理されている",
            "解法手順が明確", "反復問題が豊富",
            "Part 5と6を両方扱う", "上級語法まで扱う",
        ],
        "nice": ["短時間で演習できる", "類題付き", "復習しやすい", "時間を測れる"],
        "avoid": ["問題だけの教材", "基礎説明だけの教材", "一つの問題形式に偏る教材"],
    },
    "mock_test": {
        "label": "模試・問題演習",
        "short": "模試",
        "focus": "本番形式、模試回数、難易度、解説、時間配分",
        "needs": [
            "本番形式と時間配分を知る", "試験全体に慣れる",
            "時間内に最後まで解く", "弱点分析をしながら演習する",
            "複数回の模試で安定性を上げる", "直前期に本番感覚を維持する",
            "高得点帯のミスを減らす", "難しい模試で実力を伸ばす",
        ],
        "must": [
            "本番形式の模試を収録", "模試を複数回解ける",
            "詳しい解説がある", "スコア換算がある",
            "時間配分の助言がある", "誤答分析ができる",
        ],
        "nice": ["マークシート付き", "音声が使いやすい", "復習計画付き", "難易度が明示されている"],
        "avoid": ["解説が少ない教材", "一回分だけの教材", "古い試験形式の教材"],
    },
    "listening": {
        "label": "リスニング",
        "short": "リスニング",
        "focus": "Part 1-4、聞き取りの弱点、音声速度、音読・シャドーイング",
        "needs": [
            "英語の音に慣れる", "Part 2の応答問題を強化する",
            "音のつながりを聞き取る", "Part 3の会話を追う",
            "Part 4の説明文を聞き取る", "速い音声に対応する",
            "意図問題と細部問題を強化する", "聞き逃しを減らす",
        ],
        "must": [
            "スクリプトと訳付き", "音声変化の解説がある",
            "シャドーイングに使える", "Part 1-4を扱う",
            "速度調整ができる", "設問タイプ別に練習できる",
        ],
        "nice": ["ディクテーション付き", "短い単位で反復できる", "音声だけでも学べる"],
        "avoid": ["問題演習だけの教材", "一つのPartだけの教材", "初心者向け音声だけの教材"],
    },
    "reading": {
        "label": "リーディング・Part 7",
        "short": "Part 7",
        "focus": "Part 7、速読、時間不足、設問タイプ、解説、問題量",
        "needs": [
            "Part 7の基本的な読み方を学ぶ", "英文を前から読む",
            "シングルパッセージを安定させる", "時間不足を改善する",
            "複数文書問題に慣れる", "最後まで解き切る",
            "言い換え問題を強化する", "高難度のPart 7に対応する",
        ],
        "must": [
            "設問の根拠が分かる", "文構造の解説がある",
            "時間配分と速読法を学べる", "複数文書問題を扱う",
            "パラフレーズを詳しく扱う", "高難度問題と詳細解説がある",
        ],
        "nice": ["制限時間付き演習", "語彙解説付き", "音読にも使える", "一題ずつ学べる"],
        "avoid": ["問題数だけ多い教材", "短文中心の教材", "基礎長文だけの教材"],
    },
}

LEGACY_OUTPUT_NAMES = [
    "toeic_purchase_intents_template.csv",
    "toeic_purchase_intents_review.xlsx",
    "toeic_purchase_intents_summary.json",
]
CURRENT_OUTPUT_NAMES = [
    "toeic_query_intents_master.csv",
    "toeic_query_intents_review.xlsx",
    "toeic_query_intents_summary.json",
]
OUTPUT_COLUMNS = [
    "intent_id", "split", "category", "category_label",
    "persona_or_current_score", "target_score", "primary_need",
    "study_period", "must_have", "nice_to_have", "avoid",
    "query_focus_hint", "short_query_draft", "long_query_draft",
    "short_query_final", "long_query_final", "review_status",
    "review_note", "candidate_assignment_status", "generation_seed",
]


def validate_configuration() -> None:
    if sum(CATEGORY_QUOTAS.values()) != 80:
        raise ValueError("購入意図の合計は80件である必要があります。")
    for category, quota in CATEGORY_QUOTAS.items():
        split_total = sum(
            SPLIT_CATEGORY_QUOTAS[split][category]
            for split in ("train", "validation", "test")
        )
        if split_total != quota:
            raise ValueError(f"{category}の分割合計が不正です。")


def pick(values: list[str], index: int, offset: int = 0) -> str:
    return values[(index + offset) % len(values)]


def make_long_query(
    current: str,
    target: str,
    period: str,
    need: str,
    must: str,
    nice: str,
    avoid: str,
) -> str:
    if current == "初受験":
        opening = "TOEICは初受験です。"
    elif "・" in current:
        persona, score = current.split("・", 1)
        opening = f"{persona}で、現在は{score}です。"
    else:
        opening = f"現在は{current}です。"
    return (
        f"{opening}{period}で{target}点を目指しており、"
        f"{need}ための教材を探しています。"
        f"特に「{must}」を必須条件とし、"
        f"「{nice}」も希望しています。"
        f"「{avoid}」は避けたいです。"
    )


def build_intents(seed: int) -> pd.DataFrame:
    rng = random.Random(seed)
    rows: list[dict[str, Any]] = []

    for category, quota in CATEGORY_QUOTAS.items():
        config = CATEGORY_CONFIG[category]
        indexes = list(range(quota))
        rng.shuffle(indexes)
        split_map: dict[int, str] = {}
        cursor = 0
        for split in ("train", "validation", "test"):
            count = SPLIT_CATEGORY_QUOTAS[split][category]
            for index in indexes[cursor:cursor + count]:
                split_map[index] = split
            cursor += count

        category_offset = list(CATEGORY_QUOTAS).index(category) * 2
        for index in range(quota):
            current, target = SCORE_PATHS[
                (index + category_offset) % len(SCORE_PATHS)
            ]
            need = pick(config["needs"], index)
            must = pick(config["must"], index, 1)
            nice = pick(config["nice"], index, 2)
            avoid = pick(config["avoid"], index, 1)
            period = pick(STUDY_PERIODS, index, category_offset)
            rows.append({
                "split": split_map[index],
                "category": category,
                "category_label": config["label"],
                "persona_or_current_score": current,
                "target_score": target,
                "primary_need": need,
                "study_period": period,
                "must_have": must,
                "nice_to_have": nice,
                "avoid": avoid,
                "query_focus_hint": config["focus"],
                "short_query_draft": (
                    f"TOEIC {target}点 {config['short']} {must}"
                ),
                "long_query_draft": make_long_query(
                    current, target, period, need, must, nice, avoid
                ),
                "short_query_final": "",
                "long_query_final": "",
                "review_status": "未確認",
                "review_note": "",
                "candidate_assignment_status": "未割当",
                "generation_seed": seed,
            })

    frame = pd.DataFrame(rows)
    split_order = {"train": 0, "validation": 1, "test": 2}
    category_order = {
        category: index for index, category in enumerate(CATEGORY_QUOTAS)
    }
    frame["_split"] = frame["split"].map(split_order)
    frame["_category"] = frame["category"].map(category_order)
    frame = frame.sort_values(
        ["_split", "_category", "target_score", "primary_need"]
    ).reset_index(drop=True)
    frame["intent_id"] = [
        f"INTENT-{index:03d}" for index in range(1, len(frame) + 1)
    ]
    return frame[OUTPUT_COLUMNS]


def archive_files(
    paths: list[Path],
    archive_root: Path,
    label: str,
    preview: bool,
) -> list[tuple[Path, Path]]:
    existing = [path for path in paths if path.exists()]
    if not existing:
        return []
    archive_dir = archive_root / (
        f"{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    )
    moves = [(path, archive_dir / path.name) for path in existing]
    if not preview:
        archive_dir.mkdir(parents=True, exist_ok=True)
        for source, destination in moves:
            shutil.move(str(source), str(destination))
    return moves


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["query_intents"]
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    split_fills = {
        "train": PatternFill("solid", fgColor="E2F0D9"),
        "validation": PatternFill("solid", fgColor="FFF2CC"),
        "test": PatternFill("solid", fgColor="DDEBF7"),
    }
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row in range(2, sheet.max_row + 1):
        split = sheet.cell(row=row, column=2).value
        if split in split_fills:
            sheet.cell(row=row, column=2).fill = split_fills[split]
        for column in range(7, 19):
            sheet.cell(row=row, column=column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
    widths = [13, 12, 16, 18, 22, 12, 34, 12, 34, 30,
              30, 38, 52, 72, 52, 72, 14, 36, 16, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=index).column_letter
        ].width = width

    review_validation = DataValidation(
        type="list", formula1='"未確認,修正中,承認,除外"'
    )
    sheet.add_data_validation(review_validation)
    review_validation.add(f"Q2:Q{sheet.max_row}")
    candidate_validation = DataValidation(
        type="list", formula1='"未割当,割当中,割当済み,除外"'
    )
    sheet.add_data_validation(candidate_validation)
    candidate_validation.add(f"S2:S{sheet.max_row}")
    sheet.conditional_formatting.add(
        f"Q2:Q{sheet.max_row}",
        FormulaRule(
            formula=['Q2="承認"'],
            fill=PatternFill("solid", fgColor="C6E0B4"),
        ),
    )
    table = Table(displayName="ToeicQueryIntents", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True
    )
    sheet.add_table(table)

    for name in ("summary", "instructions"):
        target = workbook[name]
        target.freeze_panes = "A2"
        target.column_dimensions["A"].width = 28
        target.column_dimensions["B"].width = 110
        for cell in target[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in target.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def write_outputs(
    intents: pd.DataFrame,
    output_dir: Path,
    seed: int,
) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / CURRENT_OUTPUT_NAMES[0]
    xlsx_path = output_dir / CURRENT_OUTPUT_NAMES[1]
    summary_path = output_dir / CURRENT_OUTPUT_NAMES[2]

    intents.to_csv(csv_path, index=False, encoding="utf-8-sig")
    summary = {
        "design_version": "query_first_v2",
        "generation_seed": seed,
        "total_purchase_intents": int(len(intents)),
        "total_query_expressions": int(len(intents) * 2),
        "split_counts": {
            str(k): int(v)
            for k, v in intents["split"].value_counts().items()
        },
        "category_counts": {
            str(k): int(v)
            for k, v in intents["category"].value_counts().items()
        },
        "rules": [
            "商品を先に選ばず購入意図とクエリを先に固定する",
            "購入意図作成時に商品情報を参照しない",
            "Train 40・Validation 10・Test 30へ購入意図単位で分割",
            "承認後に全商品プールから候補10件を割り当てる",
        ],
        "retired_design": "対象商品80件を先に選ぶ商品先行方式",
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    instructions = pd.DataFrame([
        ["最重要", "商品名・説明・価格・URLを見ずにクエリを確定する。"],
        ["draft", "機械生成の初稿。不自然な表現だけ直し、条件を追加しない。"],
        ["final", "確定版を入力する。短文と長文は同じ購入意図を保つ。"],
        ["短文", "実際の検索語に近くし、商品名・著者名・出版社名を入れない。"],
        ["長文", "買い物AIへ相談する自然な2〜4文にする。"],
        ["Test", "JA-Adaptedプロンプトの作成・選択に使用しない。"],
        ["次工程", "全80件承認後、270件の商品プールから候補10件を割り当てる。"],
    ], columns=["項目", "ルール"])
    counts = (
        intents.groupby(["split", "category_label"])
        .size().rename("count").reset_index()
    )
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        intents.to_excel(writer, sheet_name="query_intents", index=False)
        counts.to_excel(writer, sheet_name="summary", index=False)
        instructions.to_excel(writer, sheet_name="instructions", index=False)
    style_workbook(xlsx_path)
    return csv_path, xlsx_path, summary_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "商品から独立したTOEIC購入意図80件と、"
            "短文・長文クエリのレビュー用テンプレートを作成する。"
        )
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--archive-dir", type=Path, default=DEFAULT_ARCHIVE_DIR)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--preview", action="store_true")
    args = parser.parse_args()

    validate_configuration()
    legacy_paths = [
        args.output_dir / name for name in LEGACY_OUTPUT_NAMES
    ]
    current_paths = [
        args.output_dir / name for name in CURRENT_OUTPUT_NAMES
    ]
    if any(path.exists() for path in current_paths) and not args.overwrite:
        raise FileExistsError(
            "現行出力が既にあります。再生成時は --overwrite を付けてください。"
        )

    moves = archive_files(
        legacy_paths, args.archive_dir, "旧_商品先行方式", args.preview
    )
    if args.overwrite:
        moves += archive_files(
            current_paths,
            args.archive_dir,
            "再生成前_購入意図先行方式",
            args.preview,
        )

    intents = build_intents(args.seed)
    print("旧ファイルの退避:")
    if moves:
        for source, destination in moves:
            print(f"  {source} -> {destination}")
    else:
        print("  なし")
    print(
        f"購入意図={len(intents)}件、"
        f"短文・長文={len(intents) * 2}表現"
    )
    for split in ("train", "validation", "test"):
        print(f"  {split}: {(intents['split'] == split).sum()}件")

    if args.preview:
        print("Preview only. No files were changed.")
        return

    csv_path, xlsx_path, summary_path = write_outputs(
        intents, args.output_dir, args.seed
    )
    print("購入意図先行方式の03を作成しました。")
    print(f"  {csv_path}")
    print(f"  {xlsx_path}")
    print(f"  {summary_path}")


if __name__ == "__main__":
    main()
