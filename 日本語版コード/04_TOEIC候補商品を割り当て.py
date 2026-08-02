from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_PRODUCTS = Path(
    "日本語版データ/TOEIC/02_商品プール/toeic_product_pool_final.csv"
)
DEFAULT_INTENTS = Path(
    "日本語版データ/TOEIC/03_購入意図/toeic_query_intents_review.xlsx"
)
DEFAULT_OUTPUT_DIR = Path("日本語版データ/TOEIC/04_候補商品")
DEFAULT_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
DEFAULT_TOP_K = 30
DEFAULT_BATCH_SIZE = 32

PRODUCT_REQUIRED_COLUMNS = [
    "product_id",
    "title",
    "description_clean",
]
INTENT_REQUIRED_COLUMNS = [
    "intent_id",
    "split",
    "category",
    "short_query_draft",
    "long_query_draft",
    "short_query_final",
    "long_query_final",
    "review_status",
]

# 先行研究は商品listingの title / features / description を利用する。
# TOEIC商品プールに存在する列だけを、重み付けせず1回ずつ連結する。
PRODUCT_TEXT_COLUMN_GROUPS = [
    ("title",),
    ("features", "feature_bullets", "features_clean"),
    ("description_clean", "description"),
    ("details", "details_clean"),
]


def clean_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\x00", " ").split()).strip()


def nonempty(value: Any) -> bool:
    return bool(clean_text(value))


def first_existing_column(frame: pd.DataFrame, candidates: Iterable[str]) -> str | None:
    for column in candidates:
        if column in frame.columns:
            return column
    return None


def load_products(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(
            f"商品プールが見つかりません: {path}\n"
            "02_実験用TOEIC商品プールを作成.pyを先に実行してください。"
        )
    frame = pd.read_csv(
        path, dtype=str, keep_default_na=False, encoding="utf-8-sig"
    ).fillna("")
    missing = [column for column in PRODUCT_REQUIRED_COLUMNS if column not in frame]
    if missing:
        raise ValueError("商品プールに必要な列がありません: " + ", ".join(missing))
    if frame["product_id"].duplicated().any():
        duplicated = frame.loc[
            frame["product_id"].duplicated(), "product_id"
        ].astype(str).tolist()
        raise ValueError(f"product_idが重複しています: {duplicated[:5]}")
    return frame.reset_index(drop=True)


def load_intents(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(
            f"購入意図ファイルが見つかりません: {path}\n"
            "03_TOEIC購入意図テンプレートを作成.pyを先に実行してください。"
        )
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        frame = pd.read_excel(path, sheet_name="query_intents", dtype=str).fillna("")
    else:
        frame = pd.read_csv(
            path, dtype=str, keep_default_na=False, encoding="utf-8-sig"
        ).fillna("")
    missing = [column for column in INTENT_REQUIRED_COLUMNS if column not in frame]
    if missing:
        raise ValueError("購入意図に必要な列がありません: " + ", ".join(missing))
    if frame["intent_id"].duplicated().any():
        raise ValueError("intent_idが重複しています。")
    return frame.reset_index(drop=True)


def resolve_queries(intents: pd.DataFrame, require_approved: bool) -> pd.DataFrame:
    resolved = intents.copy()
    resolved["short_query"] = resolved.apply(
        lambda row: row["short_query_final"]
        if nonempty(row["short_query_final"])
        else row["short_query_draft"],
        axis=1,
    )
    resolved["long_query"] = resolved.apply(
        lambda row: row["long_query_final"]
        if nonempty(row["long_query_final"])
        else row["long_query_draft"],
        axis=1,
    )
    resolved["query_source"] = resolved.apply(
        lambda row: (
            "final"
            if nonempty(row["short_query_final"])
            and nonempty(row["long_query_final"])
            else "draft_or_mixed"
        ),
        axis=1,
    )

    empty = resolved.loc[
        ~resolved["short_query"].map(nonempty)
        | ~resolved["long_query"].map(nonempty),
        "intent_id",
    ].astype(str).tolist()
    if empty:
        raise ValueError(f"短文または長文クエリが空です: {empty[:10]}")

    if require_approved:
        not_approved = resolved.loc[
            resolved["review_status"].ne("承認"), "intent_id"
        ].astype(str).tolist()
        missing_final = resolved.loc[
            resolved["query_source"].ne("final"), "intent_id"
        ].astype(str).tolist()
        messages: list[str] = []
        if not_approved:
            messages.append(f"未承認: {not_approved[:10]}")
        if missing_final:
            messages.append(f"final未入力: {missing_final[:10]}")
        if messages:
            raise ValueError(
                "本実験用の確定データとして使えません。\n" + "\n".join(messages)
            )
    return resolved


def resolve_product_text_columns(products: pd.DataFrame) -> list[str]:
    selected: list[str] = []
    for group in PRODUCT_TEXT_COLUMN_GROUPS:
        column = first_existing_column(products, group)
        if column is not None and column not in selected:
            selected.append(column)
    if "title" not in selected or "description_clean" not in selected:
        raise ValueError("Dense Retrievalにはtitleとdescription_cleanが必要です。")
    return selected


def combine_columns(row: pd.Series, columns: Iterable[str]) -> str:
    return "\n".join(
        value
        for column in columns
        if (value := clean_text(row.get(column, "")))
    )


def load_sentence_transformer(model_name: str, device: str | None) -> Any:
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "sentence-transformersが未導入です。\n"
            "リポジトリをpullした後、次を一度実行してください。\n"
            "  uv sync"
        ) from exc

    kwargs: dict[str, Any] = {}
    if device:
        kwargs["device"] = device
    return SentenceTransformer(model_name, **kwargs)


def encode_texts(
    model: Any,
    texts: list[str],
    batch_size: int,
    show_progress_bar: bool,
) -> np.ndarray:
    embeddings = model.encode(
        texts,
        batch_size=batch_size,
        show_progress_bar=show_progress_bar,
        convert_to_numpy=True,
        normalize_embeddings=True,
    )
    array = np.asarray(embeddings, dtype=np.float32)
    if array.ndim != 2:
        raise ValueError(f"埋め込みの形が不正です: {array.shape}")
    if not np.isfinite(array).all():
        raise ValueError("埋め込みにNaNまたは無限大が含まれています。")
    return array


def retrieve_top_k(
    product_embeddings: np.ndarray,
    query_embeddings: np.ndarray,
    top_k: int,
) -> tuple[np.ndarray, np.ndarray]:
    if product_embeddings.shape[1] != query_embeddings.shape[1]:
        raise ValueError("商品とクエリの埋め込み次元が一致しません。")
    similarity_matrix = query_embeddings @ product_embeddings.T
    top_indices = np.argsort(-similarity_matrix, axis=1, kind="stable")[:, :top_k]
    top_scores = np.take_along_axis(similarity_matrix, top_indices, axis=1)
    return top_indices, top_scores


def build_retrieval_outputs(
    products: pd.DataFrame,
    intents: pd.DataFrame,
    query_form: str,
    model_name: str,
    top_indices: np.ndarray,
    top_scores: np.ndarray,
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    query_column = f"{query_form}_query"
    rows: list[dict[str, Any]] = []
    jobs: list[dict[str, Any]] = []

    for intent_offset, (_, intent) in enumerate(
        intents.sort_values("intent_id").iterrows()
    ):
        query = clean_text(intent[query_column])
        product_items: list[dict[str, Any]] = []

        for retrieval_rank, (product_index, score) in enumerate(
            zip(top_indices[intent_offset], top_scores[intent_offset]),
            start=1,
        ):
            product = products.iloc[int(product_index)]
            row = {
                "intent_id": str(intent["intent_id"]),
                "split": str(intent["split"]),
                "category": str(intent["category"]),
                "query_form": query_form,
                "query": query,
                "query_source": str(intent["query_source"]),
                "retrieval_rank": retrieval_rank,
                "cosine_similarity": round(float(score), 8),
                "product_index": int(product_index),
                "product_id": str(product.get("product_id", "")),
                "title": str(product.get("title", "")),
                "description_clean": str(product.get("description_clean", "")),
                "author_clean": str(product.get("author_clean", "")),
                "publisher": str(product.get("publisher", "")),
                "category_hint": str(product.get("category_hint", "")),
                "score_hint": str(product.get("score_hint", "")),
                "price_yen_num": str(product.get("price_yen_num", "")),
                "product_url": str(product.get("product_url", "")),
                "embedding_model": model_name,
                "selected_top10": "",
                "selection_note": "",
            }
            rows.append(row)
            product_items.append(
                {
                    "index": retrieval_rank - 1,
                    "retrieval_rank": retrieval_rank,
                    "product_id": str(product.get("product_id", "")),
                    "title": str(product.get("title", "")),
                    "cosine_similarity": round(float(score), 8),
                }
            )

        jobs.append(
            {
                "job_type": "candidate_relevance_selection",
                "intent_id": str(intent["intent_id"]),
                "split": str(intent["split"]),
                "query_form": query_form,
                "query": query,
                "products": product_items,
                "instruction": (
                    "Select exactly 10 relevant products from the 30 titles. "
                    "Prefer smaller indices when relevance is tied. "
                    "Return a JSON array of 10 distinct integers in ascending order."
                ),
                "expected_output_example": [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
                "status": "not_executed",
            }
        )

    return pd.DataFrame(rows), jobs


def validate_retrieval(
    retrieval: pd.DataFrame,
    intent_count: int,
    top_k: int,
) -> None:
    if retrieval["intent_id"].nunique() != intent_count:
        raise ValueError("出力された購入意図数が入力と一致しません。")

    counts = retrieval.groupby("intent_id").size()
    invalid_counts = counts[counts.ne(top_k)]
    if not invalid_counts.empty:
        raise ValueError(f"上位候補数が不正です: {invalid_counts.to_dict()}")

    duplicate_pairs = retrieval.duplicated(["intent_id", "product_id"])
    if duplicate_pairs.any():
        examples = retrieval.loc[
            duplicate_pairs, ["intent_id", "product_id", "title"]
        ].head()
        raise ValueError(
            "同じ購入意図の上位30件内でproduct_idが重複しています。\n"
            + examples.to_string(index=False)
        )

    expected_ranks = list(range(1, top_k + 1))
    for intent_id, group in retrieval.groupby("intent_id", sort=False):
        if group["retrieval_rank"].astype(int).tolist() != expected_ranks:
            raise ValueError(f"{intent_id}のretrieval_rankが不正です。")

    scores = retrieval["cosine_similarity"].astype(float)
    if not np.isfinite(scores.to_numpy()).all():
        raise ValueError("コサイン類似度にNaNまたは無限大があります。")


def style_review_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    retrieval_sheet = workbook["dense_top30"]
    widths = {
        "A": 14,
        "B": 12,
        "C": 16,
        "D": 12,
        "E": 80,
        "F": 14,
        "G": 12,
        "H": 18,
        "I": 12,
        "J": 22,
        "K": 70,
        "L": 100,
        "M": 24,
        "N": 24,
        "O": 18,
        "P": 18,
        "Q": 16,
        "R": 48,
        "S": 52,
        "T": 16,
        "U": 42,
    }
    for column, width in widths.items():
        retrieval_sheet.column_dimensions[column].width = width
    table = Table(displayName="ToeicDenseTop30", ref=retrieval_sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True
    )
    retrieval_sheet.add_table(table)

    instructions = workbook["instructions"]
    instructions.column_dimensions["A"].width = 26
    instructions.column_dimensions["B"].width = 110
    workbook.save(path)


def write_outputs(
    output_dir: Path,
    retrieval: pd.DataFrame,
    jobs: list[dict[str, Any]],
    intents: pd.DataFrame,
    products: pd.DataFrame,
    product_text_columns: list[str],
    model_name: str,
    query_form: str,
    top_k: int,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    retrieval_path = output_dir / "01_dense_retrieval_top30.csv"
    jobs_path = output_dir / "02_candidate_selection_jobs.jsonl"
    summary_path = output_dir / "03_dense_retrieval_summary.json"
    review_path = output_dir / "04_dense_retrieval_review.xlsx"

    retrieval.to_csv(retrieval_path, index=False, encoding="utf-8-sig")
    jobs_path.write_text(
        "\n".join(json.dumps(job, ensure_ascii=False) for job in jobs) + "\n",
        encoding="utf-8",
    )

    digest_source = "\n".join(
        f"{row.intent_id}:{row.retrieval_rank}:{row.product_id}:{row.cosine_similarity}"
        for row in retrieval.itertuples()
    )
    summary = {
        "design_version": "toeic_egeo_dense_retrieval_v2",
        "paper_structure": (
            "dense retrieval top-30, followed by an LLM selecting 10 relevant products"
        ),
        "api_calls": 0,
        "status": "dense_top30_complete_candidate_top10_not_selected",
        "intent_count": int(retrieval["intent_id"].nunique()),
        "product_pool_count": int(len(products)),
        "retrieval_rows": int(len(retrieval)),
        "top_k": top_k,
        "query_form": query_form,
        "query_source_counts": {
            str(key): int(value)
            for key, value in intents["query_source"].value_counts().items()
        },
        "split_counts": {
            str(key): int(value)
            for key, value in intents["split"].value_counts().items()
        },
        "embedding_model": model_name,
        "embedding_normalization": "L2",
        "similarity": "cosine via normalized embedding dot product",
        "product_text_columns": product_text_columns,
        "removed_from_old_method": [
            "category pre-filter",
            "title x4 weighting",
            "target score affinity bonus",
            "avoid similarity penalty",
            "manual metadata weighting",
            "direct deterministic top-10 selection",
        ],
        "next_stage": (
            "Use an LLM relevance selector to choose exactly 10 products "
            "from each fixed top-30 list. Do not choose the target product yet."
        ),
        "retrieval_sha256": hashlib.sha256(
            digest_source.encode("utf-8")
        ).hexdigest(),
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    instructions = pd.DataFrame(
        [
            [
                "この段階の役割",
                "長文購入クエリごとに、270商品からDense Retrieval上位30件を固定する。",
            ],
            [
                "API",
                "このコードはAPIを呼ばない。02_candidate_selection_jobs.jsonlは後工程用。",
            ],
            [
                "先行研究との対応",
                "Dense Retrieval上位30件の後、LLMがタイトルから関連商品10件を選ぶ二段階方式。",
            ],
            [
                "重要",
                "現時点では候補10件も書き換え対象商品も確定していない。",
            ],
            [
                "クエリ",
                f"候補取得には{query_form}クエリを使用。同じ候補10件を後で短文実験にも共有する。",
            ],
            [
                "独自加点",
                "カテゴリ絞り込み、タイトル重み、得点帯加点、avoid減点は使用していない。",
            ],
        ],
        columns=["項目", "説明"],
    )
    with pd.ExcelWriter(review_path, engine="openpyxl") as writer:
        retrieval.to_excel(writer, sheet_name="dense_top30", index=False)
        instructions.to_excel(writer, sheet_name="instructions", index=False)
    style_review_workbook(review_path)

    return [retrieval_path, jobs_path, summary_path, review_path]


def expected_output_paths(output_dir: Path) -> list[Path]:
    return [
        output_dir / "01_dense_retrieval_top30.csv",
        output_dir / "02_candidate_selection_jobs.jsonl",
        output_dir / "03_dense_retrieval_summary.json",
        output_dir / "04_dense_retrieval_review.xlsx",
    ]


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "確定したTOEIC購入意図を使い、多言語Sentence Transformerの"
            "Dense Retrievalで商品プールから上位30件を作る。APIは呼び出さない。"
        )
    )
    parser.add_argument("--products", type=Path, default=DEFAULT_PRODUCTS)
    parser.add_argument("--intents", type=Path, default=DEFAULT_INTENTS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument(
        "--query-form",
        choices=["long", "short"],
        default="long",
        help="主実験はlong。shortは動作比較用で、正式候補の作成には使わない。",
    )
    parser.add_argument("--top-k", type=int, default=DEFAULT_TOP_K)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--device")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--require-approved", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--preview",
        action="store_true",
        help="入力と設定だけ確認し、モデル読込・埋め込み・ファイル出力を行わない。",
    )
    args = parser.parse_args()

    if args.top_k <= 0:
        raise ValueError("top-kは1以上にしてください。")
    if args.batch_size <= 0:
        raise ValueError("batch-sizeは1以上にしてください。")

    products = load_products(args.products)
    if args.top_k > len(products):
        raise ValueError(
            f"top-k={args.top_k}は商品数{len(products)}を超えています。"
        )

    all_intents = resolve_queries(load_intents(args.intents), args.require_approved)
    intents = all_intents.sort_values("intent_id").reset_index(drop=True)
    if args.limit is not None:
        if args.limit <= 0:
            raise ValueError("limitは1以上にしてください。")
        intents = intents.head(args.limit).copy()

    product_text_columns = resolve_product_text_columns(products)
    query_column = f"{args.query_form}_query"

    print("04 Dense Retrievalの入力確認が完了しました。")
    print(f"  商品プール: {len(products)}件")
    print(f"  購入意図: {len(intents)}件")
    print(f"  使用クエリ: {query_column}")
    print(f"  埋め込みモデル: {args.model}")
    print(f"  商品文章列: {', '.join(product_text_columns)}")
    print(f"  取得件数: 各{args.top_k}件")
    print("  カテゴリ絞り込み・独自加点: なし")
    print("  API呼び出し: 0回")

    if args.preview:
        print("Preview only. モデル読込とファイル出力は行っていません。")
        return

    existing = [path for path in expected_output_paths(args.output_dir) if path.exists()]
    if existing and not args.overwrite:
        raise FileExistsError(
            "出力が既にあります。再作成時は --overwrite を付けてください。\n"
            + "\n".join(str(path) for path in existing)
        )

    product_texts = [
        combine_columns(row, product_text_columns)
        for _, row in products.iterrows()
    ]
    empty_products = [
        str(products.iloc[index]["product_id"])
        for index, text in enumerate(product_texts)
        if not nonempty(text)
    ]
    if empty_products:
        raise ValueError(f"埋め込み対象の商品文章が空です: {empty_products[:10]}")

    query_texts = intents[query_column].map(clean_text).tolist()
    empty_queries = [
        str(intent_id)
        for intent_id, text in zip(intents["intent_id"], query_texts)
        if not nonempty(text)
    ]
    if empty_queries:
        raise ValueError(f"埋め込み対象クエリが空です: {empty_queries[:10]}")

    print("埋め込みモデルを読み込んでいます。初回はモデルをダウンロードします。")
    model = load_sentence_transformer(args.model, args.device)

    print("商品文章を埋め込み中です。")
    product_embeddings = encode_texts(
        model, product_texts, args.batch_size, show_progress_bar=True
    )
    print("購入クエリを埋め込み中です。")
    query_embeddings = encode_texts(
        model, query_texts, args.batch_size, show_progress_bar=True
    )

    top_indices, top_scores = retrieve_top_k(
        product_embeddings, query_embeddings, args.top_k
    )
    retrieval, jobs = build_retrieval_outputs(
        products,
        intents,
        args.query_form,
        args.model,
        top_indices,
        top_scores,
    )
    validate_retrieval(retrieval, len(intents), args.top_k)

    paths = write_outputs(
        args.output_dir,
        retrieval,
        jobs,
        intents,
        products,
        product_text_columns,
        args.model,
        args.query_form,
        args.top_k,
    )

    print("04 Dense Retrievalが完了しました。")
    print(f"  出力行数: {len(retrieval)}")
    print(f"  候補選定APIジョブ: {len(jobs)}件（未実行）")
    print("  API呼び出し: 0回")
    for path in paths:
        print(f"  {path}")


if __name__ == "__main__":
    main()
