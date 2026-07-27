from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_INPUT = Path(
    "日本語版データ/TOEIC/実験データ/toeic_product_pool_final.csv"
)
DEFAULT_OUTPUT_DIR = Path("日本語版データ/TOEIC/実験データ")
DEFAULT_SEED = 20260727

CATEGORY_QUOTAS = {
    "general": 16,
    "vocabulary": 16,
    "grammar": 14,
    "mock_test": 14,
    "listening": 10,
    "reading": 10,
}

SPLIT_CATEGORY_QUOTAS = {
    "train": {
        "general": 8,
        "vocabulary": 8,
        "grammar": 7,
        "mock_test": 7,
        "listening": 5,
        "reading": 5,
    },
    "validation": {
        "general": 2,
        "vocabulary": 2,
        "grammar": 2,
        "mock_test": 2,
        "listening": 1,
        "reading": 1,
    },
    "test": {
        "general": 6,
        "vocabulary": 6,
        "grammar": 5,
        "mock_test": 5,
        "listening": 4,
        "reading": 4,
    },
}

CATEGORY_LABELS = {
    "general": "総合対策",
    "vocabulary": "単語・熟語",
    "grammar": "文法・Part 5/6",
    "mock_test": "模試・問題演習",
    "listening": "リスニング",
    "reading": "リーディング・Part 7",
}

QUERY_FOCUS_HINTS = {
    "general": "現在スコア、目標スコア、初心者向け、一冊完結、学習期間",
    "vocabulary": "目標スコア、語彙レベル、例文、音声、反復しやすさ",
    "grammar": "Part 5/6、文法の弱点、問題演習量、解説の詳しさ",
    "mock_test": "本番形式、模試回数、難易度、解説、音声、時間配分",
    "listening": "Part 1-4、聞き取りの弱点、音声速度、シャドーイング",
    "reading": "Part 7、速読、長文読解、時間不足、問題量、解説",
}

REQUIRED_COLUMNS = [
    "product_id",
    "title",
    "author_clean",
    "price_yen_num",
    "category_hint",
    "score_hint",
    "description_clean",
    "description_length_num",
    "target_eligible_final",
    "target_ineligible_reasons",
    "quality_score",
    "review_reasons",
    "product_url",
]

OUTPUT_COLUMNS = [
    "intent_id",
    "split",
    "category",
    "category_label",
    "target_product_id",
    "target_title",
    "target_author",
    "target_score_hint",
    "target_price_yen",
    "target_description",
    "query_focus_hint",
    "short_query",
    "long_query",
    "target_reason",
    "review_status",
    "review_note",
    "selection_seed",
    "quality_score",
    "review_reasons",
    "product_url",
]


def parse_bool(value: Any) -> bool:
    return str(value).strip().casefold() in {"true", "1", "yes", "y"}


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", " ", text).strip()


def parse_score_values(value: Any) -> list[int]:
    return [
        int(x)
        for x in re.findall(
            r"(?<!\d)([2-9]\d{2}|990)(?!\d)", normalize_text(value)
        )
    ]


def score_bucket(value: Any) -> str:
    scores = parse_score_values(value)
    if not scores:
        return "unspecified"
    score = max(scores)
    if score <= 600:
        return "up_to_600"
    if score <= 800:
        return "700_to_800"
    return "900_plus"


def title_family_key(value: Any) -> str:
    """版表記と記号だけを除き、近い商品を確認しやすいキーを作る。"""
    text = normalize_text(value).casefold()
    text = re.sub(
        r"増補改訂版|完全改訂版|改訂新版|改訂版|新装版|新版", " ", text
    )
    text = re.sub(r"[\[\]【】()（）]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"[^0-9a-zぁ-んァ-ヶ一-龠]+", "", text)


def validate_quota_configuration() -> None:
    for category, total in CATEGORY_QUOTAS.items():
        split_total = sum(
            split_quota.get(category, 0)
            for split_quota in SPLIT_CATEGORY_QUOTAS.values()
        )
        if split_total != total:
            raise ValueError(
                f"カテゴリ {category} の総数が一致しません: "
                f"category={total}, split_total={split_total}"
            )

    split_totals = {
        split: sum(category_counts.values())
        for split, category_counts in SPLIT_CATEGORY_QUOTAS.items()
    }
    if split_totals != {"train": 40, "validation": 10, "test": 30}:
        raise ValueError(f"分割件数が想定外です: {split_totals}")


def prepare_candidates(df: pd.DataFrame, seed: int) -> pd.DataFrame:
    work = df.copy().fillna("")
    work["target_eligible_bool"] = work["target_eligible_final"].map(
        parse_bool
    )
    work["quality_score_num"] = pd.to_numeric(
        work["quality_score"], errors="coerce"
    ).fillna(0.0)
    work["description_length_num2"] = pd.to_numeric(
        work["description_length_num"], errors="coerce"
    ).fillna(0).astype(int)
    work["score_bucket"] = work["score_hint"].map(score_bucket)
    work["family_key"] = work["title"].map(title_family_key)

    work = work.loc[
        work["target_eligible_bool"]
        & work["category_hint"].isin(CATEGORY_QUOTAS)
    ].copy()

    rng = np.random.default_rng(seed)
    work["selection_random"] = rng.random(len(work))
    return work


def select_targets(candidates: pd.DataFrame) -> pd.DataFrame:
    selected_parts: list[pd.DataFrame] = []

    for category, quota in CATEGORY_QUOTAS.items():
        category_rows = candidates.loc[
            candidates["category_hint"].eq(category)
        ].copy()
        if len(category_rows) < quota:
            raise ValueError(
                f"対象商品候補が不足しています: {category}="
                f"{len(category_rows)}件、必要={quota}件"
            )

        # target_eligible_final を通過した商品から固定シードで抽出する。
        # quality_scoreは抽出条件にはせず、選定後の確認情報として残す。
        category_rows = category_rows.sort_values(
            ["selection_random", "product_id"],
            ascending=[True, True],
        )
        selected_parts.append(category_rows.head(quota))

    selected = pd.concat(selected_parts, ignore_index=True)
    if selected["product_id"].duplicated().any():
        duplicates = selected.loc[
            selected["product_id"].duplicated(keep=False), "product_id"
        ].tolist()
        raise ValueError(f"対象商品IDが重複しています: {duplicates}")
    return selected


def assign_splits(selected: pd.DataFrame, seed: int) -> pd.DataFrame:
    rng = np.random.default_rng(seed + 1)
    assigned_parts: list[pd.DataFrame] = []

    for category in CATEGORY_QUOTAS:
        category_rows = selected.loc[
            selected["category_hint"].eq(category)
        ].copy()
        category_rows["split_random"] = rng.random(len(category_rows))
        category_rows = category_rows.sort_values(
            ["split_random", "product_id"],
            ascending=[True, True],
        ).reset_index(drop=True)

        start = 0
        for split in ["train", "validation", "test"]:
            count = SPLIT_CATEGORY_QUOTAS[split][category]
            part = category_rows.iloc[start : start + count].copy()
            part["split"] = split
            assigned_parts.append(part)
            start += count

        if start != len(category_rows):
            raise ValueError(
                f"分割後の件数が一致しません: {category}="
                f"{start}/{len(category_rows)}"
            )

    assigned = pd.concat(assigned_parts, ignore_index=True)
    split_order = {"train": 0, "validation": 1, "test": 2}
    category_order = {
        category: i for i, category in enumerate(CATEGORY_QUOTAS)
    }
    assigned["_split_order"] = assigned["split"].map(split_order)
    assigned["_category_order"] = assigned["category_hint"].map(
        category_order
    )
    assigned = assigned.sort_values(
        ["_split_order", "_category_order", "split_random", "product_id"]
    ).reset_index(drop=True)
    assigned["intent_id"] = [
        f"INTENT-{i:03d}" for i in range(1, len(assigned) + 1)
    ]
    return assigned.drop(columns=["_split_order", "_category_order"])


def make_template(assigned: pd.DataFrame, seed: int) -> pd.DataFrame:
    template = pd.DataFrame(
        {
            "intent_id": assigned["intent_id"],
            "split": assigned["split"],
            "category": assigned["category_hint"],
            "category_label": assigned["category_hint"].map(
                CATEGORY_LABELS
            ),
            "target_product_id": assigned["product_id"],
            "target_title": assigned["title"],
            "target_author": assigned["author_clean"],
            "target_score_hint": assigned["score_hint"],
            "target_price_yen": pd.to_numeric(
                assigned["price_yen_num"], errors="coerce"
            ).fillna(0).astype(int),
            "target_description": assigned["description_clean"],
            "query_focus_hint": assigned["category_hint"].map(
                QUERY_FOCUS_HINTS
            ),
            "short_query": "",
            "long_query": "",
            "target_reason": "",
            "review_status": "未確認",
            "review_note": "",
            "selection_seed": seed,
            "quality_score": pd.to_numeric(
                assigned["quality_score"], errors="coerce"
            ).fillna(0.0),
            "review_reasons": assigned["review_reasons"],
            "product_url": assigned["product_url"],
        }
    )
    return template[OUTPUT_COLUMNS]


def build_summary(
    template: pd.DataFrame, assigned: pd.DataFrame, seed: int
) -> dict[str, Any]:
    category_counts = (
        template["category"]
        .value_counts()
        .sort_index()
        .astype(int)
        .to_dict()
    )
    split_counts = template["split"].value_counts().astype(int).to_dict()
    split_category_counts = {
        split: (
            template.loc[template["split"].eq(split), "category"]
            .value_counts()
            .sort_index()
            .astype(int)
            .to_dict()
        )
        for split in ["train", "validation", "test"]
    }
    score_bucket_counts = (
        assigned["score_bucket"]
        .value_counts()
        .sort_index()
        .astype(int)
        .to_dict()
    )
    repeated_family_count = int(
        (
            assigned["family_key"].ne("")
            & assigned["family_key"].duplicated(keep=False)
        ).sum()
    )

    return {
        "input_file": str(DEFAULT_INPUT),
        "selection_seed": seed,
        "selected_target_rows": int(len(template)),
        "split_counts": split_counts,
        "category_counts": category_counts,
        "split_category_counts": split_category_counts,
        "score_bucket_counts": score_bucket_counts,
        "rows_in_repeated_title_families": repeated_family_count,
        "rules": [
            "target_eligible_final=Trueの商品だけを対象商品候補に使用",
            "カテゴリ別に固定件数を抽出",
            "quality_scoreによる上位選抜は行わず、固定乱数シードで抽出",
            "Train 40・Validation 10・Test 30へカテゴリ比率を保って分割",
            "80件の対象商品IDは重複させない",
            "短文・長文クエリと採用理由は次工程で記入・確認",
        ],
    }


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb["purchase_intents"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    split_fills = {
        "train": PatternFill("solid", fgColor="E2F0D9"),
        "validation": PatternFill("solid", fgColor="FFF2CC"),
        "test": PatternFill("solid", fgColor="DDEBF7"),
    }

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, ws.max_row + 1):
        split_value = ws.cell(row=row, column=2).value
        if split_value in split_fills:
            ws.cell(row=row, column=2).fill = split_fills[split_value]
        for col in [6, 10, 11, 12, 13, 14, 16]:
            ws.cell(row=row, column=col).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    widths = {
        "A": 13,
        "B": 12,
        "C": 16,
        "D": 18,
        "E": 18,
        "F": 44,
        "G": 18,
        "H": 16,
        "I": 14,
        "J": 55,
        "K": 36,
        "L": 30,
        "M": 55,
        "N": 42,
        "O": 14,
        "P": 30,
        "Q": 14,
        "R": 14,
        "S": 28,
        "T": 38,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 28

    status_validation = DataValidation(
        type="list",
        formula1='"未確認,修正中,承認,除外"',
        allow_blank=False,
    )
    ws.add_data_validation(status_validation)
    status_validation.add(f"O2:O{ws.max_row}")

    ws.conditional_formatting.add(
        f"O2:O{ws.max_row}",
        FormulaRule(
            formula=['O2="承認"'],
            fill=PatternFill("solid", fgColor="C6E0B4"),
        ),
    )
    ws.conditional_formatting.add(
        f"O2:O{ws.max_row}",
        FormulaRule(
            formula=['O2="除外"'],
            fill=PatternFill("solid", fgColor="F4CCCC"),
        ),
    )

    table = Table(displayName="PurchaseIntentsTable", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    instructions = wb["instructions"]
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 100
    instructions.freeze_panes = "A2"
    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in instructions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = wb["summary"]
    summary.column_dimensions["A"].width = 22
    for col in range(2, summary.max_column + 1):
        summary.column_dimensions[get_column_letter(col)].width = 15
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    wb.save(path)


def write_outputs(
    template: pd.DataFrame,
    assigned: pd.DataFrame,
    summary: dict[str, Any],
    output_dir: Path,
) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "toeic_purchase_intents_template.csv"
    xlsx_path = output_dir / "toeic_purchase_intents_review.xlsx"
    summary_path = output_dir / "toeic_purchase_intents_summary.json"

    template.to_csv(csv_path, index=False, encoding="utf-8-sig")
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    category_summary = (
        template.pivot_table(
            index="category_label",
            columns="split",
            values="intent_id",
            aggfunc="count",
            fill_value=0,
        )
        .reindex([CATEGORY_LABELS[k] for k in CATEGORY_QUOTAS])
        .reset_index()
    )
    for split in ["train", "validation", "test"]:
        if split not in category_summary.columns:
            category_summary[split] = 0
    category_summary["total"] = category_summary[
        ["train", "validation", "test"]
    ].sum(axis=1)

    instructions = pd.DataFrame(
        [
            [
                "short_query",
                "3〜8語程度を目安に、実際の検索キーワードとして自然に書く。対象商品の固有名詞は原則入れない。",
            ],
            [
                "long_query",
                "同じ購入意図を2〜3文で具体化する。現在スコア、目標、弱点、教材形式など必要な条件だけを書く。",
            ],
            [
                "target_reason",
                "候補商品の説明に基づき、この商品が購入意図に最も合う理由を1〜2文で記録する。",
            ],
            [
                "意味対応",
                "short_queryとlong_queryは表現の長さだけを変え、必要条件や正解対象商品を変えない。",
            ],
            [
                "禁止事項",
                "商品説明にない特典・対象スコア・実績を推測で追加しない。商品名をそのまま検索クエリにしない。",
            ],
            [
                "review_status",
                "未確認→修正中→承認の順で更新する。不適切な対象商品は除外とし、review_noteへ理由を書く。",
            ],
            [
                "分割",
                "Testの内容はJA-Adaptedプロンプトの作成・選択に使用しない。",
            ],
        ],
        columns=["項目", "作業ルール"],
    )

    selected_details = assigned[
        [
            "intent_id",
            "split",
            "category_hint",
            "score_bucket",
            "family_key",
            "description_length_num2",
            "selection_random",
            "split_random",
            "product_id",
            "title",
        ]
    ].copy()

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        template.to_excel(writer, sheet_name="purchase_intents", index=False)
        category_summary.to_excel(writer, sheet_name="summary", index=False)
        instructions.to_excel(writer, sheet_name="instructions", index=False)
        selected_details.to_excel(writer, sheet_name="selection_log", index=False)

    style_workbook(xlsx_path)
    return csv_path, xlsx_path, summary_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "最終TOEIC商品プールから対象商品80件を抽出し、"
            "購入意図・短文・長文クエリを記入するテンプレートを作る。"
        )
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    args = parser.parse_args()

    validate_quota_configuration()
    if not args.input.exists():
        raise FileNotFoundError(
            f"入力ファイルが見つかりません: {args.input}\n"
            "09_TOEIC商品プール最終化.pyを先に実行してください。"
        )

    raw = pd.read_csv(
        args.input,
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
    )
    missing = [column for column in REQUIRED_COLUMNS if column not in raw.columns]
    if missing:
        raise ValueError("必要な列がありません: " + ", ".join(missing))

    candidates = prepare_candidates(raw, seed=args.seed)
    selected = select_targets(candidates)
    assigned = assign_splits(selected, seed=args.seed)
    template = make_template(assigned, seed=args.seed)
    summary = build_summary(template, assigned, seed=args.seed)
    csv_path, xlsx_path, summary_path = write_outputs(
        template=template,
        assigned=assigned,
        summary=summary,
        output_dir=args.output_dir,
    )

    print("TOEIC購入意図テンプレートの作成が完了しました。")
    print(f"  対象商品: {len(template)}件")
    print("  分割:")
    for split in ["train", "validation", "test"]:
        print(f"    {split}: {(template['split'] == split).sum()}件")
    print("  カテゴリ:")
    for category in CATEGORY_QUOTAS:
        count = int((template["category"] == category).sum())
        print(f"    {category}: {count}件")
    print(f"  出力: {csv_path}")
    print(f"        {xlsx_path}")
    print(f"        {summary_path}")
    print(
        "  次はreview.xlsxのshort_query、long_query、target_reasonを作成します。"
    )


if __name__ == "__main__":
    main()
