from __future__ import annotations

import argparse
import hashlib
import json
import random
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_RETRIEVAL = Path(
    "日本語版データ/TOEIC/04_候補商品/01_dense_retrieval_top30.csv"
)
DEFAULT_JOBS = Path(
    "日本語版データ/TOEIC/04_候補商品/02_candidate_selection_jobs.jsonl"
)
DEFAULT_INTENTS = Path(
    "日本語版データ/TOEIC/03_購入意図/toeic_query_intents_review.xlsx"
)
DEFAULT_CONFIG = Path(
    "日本語版設定/E_GEO先行研究_候補10件選定設定.json"
)
DEFAULT_OUTPUT_DIR = Path("日本語版データ/TOEIC/04_候補商品")
DEFAULT_MODEL = "gpt-5-mini"
DEFAULT_SEED = 42
SPLITS = ("train", "validation", "test")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def digest(value: Any) -> str:
    raw = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def clean_text(value: Any) -> str:
    return " ".join(str(value or "").split()).strip()


def read_json(path: Path) -> Any:
    if not path.exists():
        raise FileNotFoundError(f"ファイルが見つかりません: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"JSONLが見つかりません: {path}")
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(
        path.read_text(encoding="utf-8").splitlines(),
        start=1,
    ):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(
                f"JSONLを読めません: {path}:{line_number}"
            ) from exc
        if not isinstance(value, dict):
            raise ValueError(f"JSONLの各行はobjectである必要があります: {line_number}")
        rows.append(value)
    return rows


def read_result_cache(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    cache: dict[str, dict[str, Any]] = {}
    for row in read_jsonl_rows(path):
        job_id = str(row.get("job_id", ""))
        if not job_id:
            raise ValueError(f"結果JSONLにjob_idがありません: {path}")
        cache[job_id] = row
    return cache


def append_jsonl(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(row, ensure_ascii=False) + "\n")


def load_intents(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"購入意図ファイルが見つかりません: {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        frame = pd.read_excel(path, sheet_name="query_intents", dtype=str).fillna("")
    else:
        frame = pd.read_csv(
            path,
            dtype=str,
            keep_default_na=False,
            encoding="utf-8-sig",
        ).fillna("")

    required = [
        "intent_id",
        "split",
        "short_query_draft",
        "long_query_draft",
        "short_query_final",
        "long_query_final",
    ]
    missing = [column for column in required if column not in frame]
    if missing:
        raise ValueError("購入意図に必要な列がありません: " + ", ".join(missing))
    if frame["intent_id"].duplicated().any():
        raise ValueError("購入意図のintent_idが重複しています。")

    frame["short_query"] = frame.apply(
        lambda row: clean_text(row["short_query_final"])
        or clean_text(row["short_query_draft"]),
        axis=1,
    )
    frame["long_query"] = frame.apply(
        lambda row: clean_text(row["long_query_final"])
        or clean_text(row["long_query_draft"]),
        axis=1,
    )
    if frame["short_query"].eq("").any() or frame["long_query"].eq("").any():
        raise ValueError("短文または長文クエリが空の購入意図があります。")
    return frame.reset_index(drop=True)


def load_retrieval(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Dense Retrieval結果が見つかりません: {path}")
    frame = pd.read_csv(
        path,
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
    ).fillna("")
    required = [
        "intent_id",
        "split",
        "query_form",
        "query",
        "retrieval_rank",
        "product_id",
        "title",
        "description_clean",
    ]
    missing = [column for column in required if column not in frame]
    if missing:
        raise ValueError(
            "Dense Retrieval結果に必要な列がありません: " + ", ".join(missing)
        )
    frame["retrieval_rank"] = frame["retrieval_rank"].astype(int)
    frame = frame.sort_values(
        ["intent_id", "retrieval_rank"],
        kind="stable",
    ).reset_index(drop=True)
    return frame


def validate_inputs(
    retrieval: pd.DataFrame,
    jobs: list[dict[str, Any]],
) -> None:
    intent_ids = sorted(retrieval["intent_id"].astype(str).unique().tolist())
    if not intent_ids:
        raise ValueError("Dense Retrieval結果が空です。")

    counts = retrieval.groupby("intent_id").size()
    invalid_counts = counts[counts.ne(30)]
    if not invalid_counts.empty:
        raise ValueError(f"上位30件でない購入意図があります: {invalid_counts.to_dict()}")

    if retrieval.duplicated(["intent_id", "product_id"]).any():
        raise ValueError("同じ購入意図の上位30件内で商品が重複しています。")

    expected_ranks = list(range(1, 31))
    for intent_id, group in retrieval.groupby("intent_id", sort=False):
        ranks = group["retrieval_rank"].astype(int).tolist()
        if ranks != expected_ranks:
            raise ValueError(f"{intent_id}: retrieval_rankが1〜30ではありません。")

    if len(jobs) != len(intent_ids):
        raise ValueError(
            f"候補選定ジョブ数{len(jobs)}が購入意図数{len(intent_ids)}と一致しません。"
        )

    seen_jobs: set[str] = set()
    for job in jobs:
        intent_id = str(job.get("intent_id", ""))
        if intent_id in seen_jobs:
            raise ValueError(f"候補選定ジョブのintent_idが重複しています: {intent_id}")
        seen_jobs.add(intent_id)
        if intent_id not in intent_ids:
            raise ValueError(f"ジョブに未知のintent_idがあります: {intent_id}")
        products = job.get("products", [])
        if not isinstance(products, list) or len(products) != 30:
            raise ValueError(f"{intent_id}: ジョブの商品数が30件ではありません。")
        indexes = [int(product["index"]) for product in products]
        if indexes != list(range(30)):
            raise ValueError(f"{intent_id}: ジョブindexが0〜29ではありません。")


def format_products(job: dict[str, Any]) -> str:
    return "\n".join(
        f"{int(product['index'])}: {clean_text(product.get('title', ''))}"
        for product in job["products"]
    )


def build_api_jobs(
    jobs: list[dict[str, Any]],
    config: dict[str, Any],
    model: str,
) -> list[dict[str, Any]]:
    system_prompt = config["system_prompt"]["source_prompt_en"]
    user_template = config["user_prompt"]["source_prompt_en"]
    output: list[dict[str, Any]] = []

    for source_job in jobs:
        input_prompt = user_template.format(
            query=source_job["query"],
            products_text=format_products(source_job),
        )
        payload = {
            "model": model,
            "instructions": system_prompt,
            "input": input_prompt,
            "store": False,
        }
        stable_id = (
            f"candidate_select__{source_job['intent_id']}__{model}__"
            f"{digest(payload)[:12]}"
        )
        output.append(
            {
                "job_id": stable_id,
                "stage": "candidate_select",
                "intent_id": str(source_job["intent_id"]),
                "split": str(source_job["split"]),
                "query_form": str(source_job["query_form"]),
                "model": model,
                "payload": payload,
            }
        )
    return output


def output_text(response: Any) -> str:
    text = getattr(response, "output_text", None)
    if text:
        return str(text)
    raw = response.model_dump() if hasattr(response, "model_dump") else response
    return "".join(
        content.get("text", "")
        for item in raw.get("output", [])
        for content in item.get("content", [])
        if content.get("type") == "output_text"
    )


def usage(response: Any) -> dict[str, int]:
    item = getattr(response, "usage", None)
    if item is None:
        return {
            "input_tokens": 0,
            "output_tokens": 0,
            "total_tokens": 0,
        }
    return {
        "input_tokens": int(getattr(item, "input_tokens", 0) or 0),
        "output_tokens": int(getattr(item, "output_tokens", 0) or 0),
        "total_tokens": int(getattr(item, "total_tokens", 0) or 0),
    }


def api_call(
    client: Any,
    payload: dict[str, Any],
    retries: int,
) -> tuple[str, Any]:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            response = client.responses.create(**payload)
            text = output_text(response).strip()
            if not text:
                raise ValueError("API応答が空です。")
            return text, response
        except Exception as exc:
            last_error = exc
            if attempt + 1 < retries:
                time.sleep(min(2**attempt + random.random(), 8))
    raise RuntimeError(f"API呼び出しに失敗しました: {last_error}") from last_error


def parse_selection(text: str) -> list[int]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`").strip()
        if cleaned.startswith("json"):
            cleaned = cleaned[4:].strip()
    value = json.loads(cleaned)
    if not isinstance(value, list):
        raise ValueError("候補選定結果はJSON配列である必要があります。")
    selected = [int(number) for number in value]
    if len(selected) != 10:
        raise ValueError(f"選定件数が10件ではありません: {selected}")
    if len(set(selected)) != 10:
        raise ValueError(f"選定indexが重複しています: {selected}")
    if any(number < 0 or number > 29 for number in selected):
        raise ValueError(f"選定indexが0〜29の範囲外です: {selected}")
    if selected != sorted(selected):
        raise ValueError(f"選定indexが昇順ではありません: {selected}")
    return selected


def execute_jobs(
    api_jobs: list[dict[str, Any]],
    cache: dict[str, dict[str, Any]],
    result_path: Path,
    retries: int,
) -> None:
    from openai import OpenAI

    client = OpenAI()
    for number, job in enumerate(api_jobs, start=1):
        existing = cache.get(job["job_id"])
        if existing and existing.get("status") == "success":
            continue

        print(
            f"[candidate_select {number}/{len(api_jobs)}] "
            f"{job['intent_id']} {job['model']}"
        )
        started = now()
        try:
            text, response = api_call(client, job["payload"], retries)
            selected = parse_selection(text)
            row = {
                key: value
                for key, value in job.items()
                if key != "payload"
            }
            row.update(
                {
                    "status": "success",
                    "selected_indices": selected,
                    "raw_output": text,
                    "usage": usage(response),
                    "response_id": getattr(response, "id", None),
                    "started_at": started,
                    "finished_at": now(),
                }
            )
        except Exception as exc:
            row = {
                key: value
                for key, value in job.items()
                if key != "payload"
            }
            row.update(
                {
                    "status": "failed",
                    "error": str(exc),
                    "started_at": started,
                    "finished_at": now(),
                }
            )

        append_jsonl(result_path, row)
        cache[job["job_id"]] = row


def stable_target_positions(
    all_intent_ids: list[str],
    seed: int,
) -> dict[str, int]:
    ordered_ids = sorted(all_intent_ids)
    rng = np.random.default_rng(seed)
    positions = rng.integers(1, 11, size=len(ordered_ids))
    return {
        intent_id: int(position)
        for intent_id, position in zip(ordered_ids, positions)
    }


def selected_result_by_intent(
    api_jobs: list[dict[str, Any]],
    cache: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    missing: list[str] = []
    for job in api_jobs:
        result = cache.get(job["job_id"])
        if not result or result.get("status") != "success":
            missing.append(str(job["intent_id"]))
            continue
        parse_selection(json.dumps(result["selected_indices"]))
        output[str(job["intent_id"])] = result
    if missing:
        raise ValueError(
            "候補10件を確定できない購入意図があります: "
            + ", ".join(missing[:10])
        )
    return output


def build_final_outputs(
    retrieval: pd.DataFrame,
    intents: pd.DataFrame,
    api_jobs: list[dict[str, Any]],
    cache: dict[str, dict[str, Any]],
    seed: int,
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    result_map = selected_result_by_intent(api_jobs, cache)
    all_intent_ids = sorted(retrieval["intent_id"].astype(str).unique().tolist())
    target_positions = stable_target_positions(all_intent_ids, seed)
    intent_map = {
        str(row.intent_id): row
        for row in intents.itertuples(index=False)
    }

    assignment_rows: list[dict[str, Any]] = []
    instances: list[dict[str, Any]] = []

    for intent_id in sorted(result_map):
        if intent_id not in intent_map:
            raise ValueError(f"購入意図ファイルにintent_idがありません: {intent_id}")
        intent = intent_map[intent_id]
        group = retrieval.loc[
            retrieval["intent_id"].astype(str).eq(intent_id)
        ].sort_values("retrieval_rank")
        row_by_index = {
            int(row.retrieval_rank) - 1: row
            for row in group.itertuples(index=False)
        }
        selected_indices = [
            int(number)
            for number in result_map[intent_id]["selected_indices"]
        ]
        target_position = target_positions[intent_id]
        product_payloads: list[dict[str, Any]] = []

        for candidate_position, selected_index in enumerate(
            selected_indices,
            start=1,
        ):
            source = row_by_index[selected_index]
            is_target = candidate_position == target_position
            assignment = {
                "intent_id": intent_id,
                "split": str(intent.split),
                "category": str(getattr(intent, "category", "")),
                "short_query": str(intent.short_query),
                "long_query": str(intent.long_query),
                "candidate_position": candidate_position,
                "source_top30_index": selected_index,
                "retrieval_rank": int(source.retrieval_rank),
                "cosine_similarity": float(source.cosine_similarity),
                "target_candidate_position": target_position,
                "is_target": is_target,
                "product_id": str(source.product_id),
                "title": str(source.title),
                "description_clean": str(source.description_clean),
                "author_clean": str(getattr(source, "author_clean", "")),
                "publisher": str(getattr(source, "publisher", "")),
                "category_hint": str(getattr(source, "category_hint", "")),
                "score_hint": str(getattr(source, "score_hint", "")),
                "price_yen_num": str(getattr(source, "price_yen_num", "")),
                "product_url": str(getattr(source, "product_url", "")),
                "selection_model": str(result_map[intent_id]["model"]),
                "target_selection_seed": seed,
                "target_selection_rule": "fixed_seed_uniform_index_1_to_10",
            }
            assignment_rows.append(assignment)
            product_payloads.append(
                {
                    "candidate_position": candidate_position,
                    "product_id": str(source.product_id),
                    "title": str(source.title),
                    "author": str(getattr(source, "author_clean", "")),
                    "price_yen": str(getattr(source, "price_yen_num", "")),
                    "category": str(getattr(source, "category_hint", "")),
                    "description": str(source.description_clean),
                    "is_target": is_target,
                    "retrieval_rank": int(source.retrieval_rank),
                    "cosine_similarity": float(source.cosine_similarity),
                }
            )

        target_product = product_payloads[target_position - 1]
        instances.append(
            {
                "intent_id": intent_id,
                "split": str(intent.split),
                "category": str(getattr(intent, "category", "")),
                "queries": {
                    "short": str(intent.short_query),
                    "long": str(intent.long_query),
                    "source": "final",
                },
                "target_candidate_position": target_position,
                "target_product_id": target_product["product_id"],
                "products": product_payloads,
            }
        )

    assignments = pd.DataFrame(assignment_rows)
    validate_final_assignments(assignments)
    return assignments, instances


def validate_final_assignments(assignments: pd.DataFrame) -> None:
    if assignments.empty:
        raise ValueError("候補10件の最終割当が空です。")
    counts = assignments.groupby("intent_id").size()
    if not counts.eq(10).all():
        raise ValueError(
            f"候補件数が10件でない購入意図があります: "
            f"{counts[counts.ne(10)].to_dict()}"
        )
    if assignments.duplicated(["intent_id", "product_id"]).any():
        raise ValueError("最終候補10件内で商品が重複しています。")
    target_counts = assignments.groupby("intent_id")["is_target"].sum()
    if not target_counts.eq(1).all():
        raise ValueError(
            f"対象商品数が1件でない購入意図があります: "
            f"{target_counts[target_counts.ne(1)].to_dict()}"
        )


def style_review_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    assignments = workbook["candidate_assignments"]
    widths = {
        "A": 14,
        "B": 12,
        "C": 16,
        "D": 45,
        "E": 80,
        "F": 12,
        "G": 14,
        "H": 14,
        "I": 12,
        "J": 12,
        "K": 18,
        "L": 16,
        "M": 70,
        "N": 100,
        "O": 24,
        "P": 24,
        "Q": 18,
        "R": 18,
        "S": 16,
        "T": 48,
        "U": 30,
        "V": 18,
        "W": 34,
    }
    for column, width in widths.items():
        assignments.column_dimensions[column].width = width

    table = Table(
        displayName="ToeicCandidateAssignments",
        ref=assignments.dimensions,
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
    )
    assignments.add_table(table)

    instructions = workbook["instructions"]
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 110
    workbook.save(path)


def write_plan(
    output_dir: Path,
    retrieval: pd.DataFrame,
    api_jobs: list[dict[str, Any]],
    config: dict[str, Any],
    model: str,
    splits: list[str],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "05_candidate_selection_plan.json"
    plan = {
        "created_at": now(),
        "stage": "candidate_selection_plan",
        "api_calls_executed": 0,
        "api_calls_planned": len(api_jobs),
        "intent_count": len(api_jobs),
        "splits": splits,
        "model": model,
        "input_candidate_count": 30,
        "output_candidate_count": 10,
        "retrieval_rows": int(len(retrieval)),
        "prompt_source": config["source"],
        "target_product_selected": False,
        "next_step": (
            "After API use is explicitly approved, execute the candidate "
            "selection jobs. Then choose one target position per final set "
            "with fixed seed 42."
        ),
        "job_ids": [job["job_id"] for job in api_jobs],
    }
    path.write_text(
        json.dumps(plan, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return path


def write_final_outputs(
    output_dir: Path,
    assignments: pd.DataFrame,
    instances: list[dict[str, Any]],
    cache: dict[str, dict[str, Any]],
    seed: int,
    model: str,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    assignments_path = output_dir / "07_candidate_assignments.csv"
    instances_path = output_dir / "08_toeic_experiment_instances.json"
    summary_path = output_dir / "09_candidate_selection_summary.json"
    review_path = output_dir / "10_candidate_assignments_review.xlsx"

    assignments.to_csv(
        assignments_path,
        index=False,
        encoding="utf-8-sig",
    )
    instances_path.write_text(
        json.dumps(instances, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    successful_rows = [
        row
        for row in cache.values()
        if row.get("status") == "success"
    ]
    token_totals = {
        "input_tokens": sum(
            int(row.get("usage", {}).get("input_tokens", 0) or 0)
            for row in successful_rows
        ),
        "output_tokens": sum(
            int(row.get("usage", {}).get("output_tokens", 0) or 0)
            for row in successful_rows
        ),
        "total_tokens": sum(
            int(row.get("usage", {}).get("total_tokens", 0) or 0)
            for row in successful_rows
        ),
    }
    digest_source = "\n".join(
        f"{row.intent_id}:{row.candidate_position}:{row.product_id}:"
        f"{int(row.is_target)}"
        for row in assignments.itertuples()
    )
    summary = {
        "design_version": "toeic_egeo_candidate_selection_v2",
        "status": "candidate_top10_and_target_complete",
        "intent_count": int(assignments["intent_id"].nunique()),
        "candidate_rows": int(len(assignments)),
        "candidate_count_per_intent": 10,
        "selection_model": model,
        "selection_rule": (
            "GPT relevance selection from fixed dense top-30 titles; "
            "ties prefer smaller retrieval indices"
        ),
        "target_selection_seed": seed,
        "target_selection_rule": "fixed_seed_uniform_index_1_to_10",
        "token_usage": token_totals,
        "assignment_sha256": hashlib.sha256(
            digest_source.encode("utf-8")
        ).hexdigest(),
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    instructions = pd.DataFrame(
        [
            [
                "候補10件",
                "Dense Retrieval上位30件から、先行研究と同じ形式のLLM関連性判定で10件を選定した。",
            ],
            [
                "候補順",
                "LLM出力indexは昇順であり、Dense Retrieval順位の早い順を候補1〜10にした。",
            ],
            [
                "対象商品",
                f"候補10件の確定後、seed {seed}で各購入意図の対象位置を一度だけ決めた。",
            ],
            [
                "短文・長文",
                "主実験の長文で候補を作り、同じ10件と対象商品を短文追加実験にも共有する。",
            ],
            [
                "禁止事項",
                "初期LLM順位やリライト結果を見た後で候補または対象商品を選び直さない。",
            ],
        ],
        columns=["項目", "説明"],
    )
    with pd.ExcelWriter(review_path, engine="openpyxl") as writer:
        assignments.to_excel(
            writer,
            sheet_name="candidate_assignments",
            index=False,
        )
        instructions.to_excel(
            writer,
            sheet_name="instructions",
            index=False,
        )
    style_review_workbook(review_path)

    return [
        assignments_path,
        instances_path,
        summary_path,
        review_path,
    ]


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Dense Retrieval上位30件から、先行研究と同じLLM関連性判定で "
            "候補10件を選ぶ。既定では計画だけ作成し、APIは呼び出さない。"
        )
    )
    parser.add_argument("--retrieval", type=Path, default=DEFAULT_RETRIEVAL)
    parser.add_argument("--jobs", type=Path, default=DEFAULT_JOBS)
    parser.add_argument("--intents", type=Path, default=DEFAULT_INTENTS)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument(
        "--splits",
        nargs="+",
        choices=SPLITS,
        default=list(SPLITS),
    )
    parser.add_argument("--limit", type=int)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--max-retries", type=int, default=3)
    parser.add_argument(
        "--rebuild-final",
        action="store_true",
        help="既存の成功結果から候補10件と対象商品ファイルを再生成する。",
    )
    args = parser.parse_args()

    if args.limit is not None and args.limit <= 0:
        raise ValueError("limitは1以上にしてください。")
    if args.max_retries <= 0:
        raise ValueError("max-retriesは1以上にしてください。")

    load_dotenv()
    retrieval_all = load_retrieval(args.retrieval)
    source_jobs_all = read_jsonl_rows(args.jobs)
    intents = load_intents(args.intents)
    config = read_json(args.config)
    validate_inputs(retrieval_all, source_jobs_all)

    split_set = set(args.splits)
    allowed_intent_ids = set(
        retrieval_all.loc[
            retrieval_all["split"].isin(split_set),
            "intent_id",
        ].astype(str)
    )
    source_jobs = [
        job
        for job in source_jobs_all
        if str(job["intent_id"]) in allowed_intent_ids
    ]
    source_jobs.sort(key=lambda row: str(row["intent_id"]))
    if args.limit is not None:
        source_jobs = source_jobs[: args.limit]

    selected_ids = {
        str(job["intent_id"])
        for job in source_jobs
    }
    retrieval = retrieval_all.loc[
        retrieval_all["intent_id"].astype(str).isin(selected_ids)
    ].copy()
    api_jobs = build_api_jobs(source_jobs, config, args.model)

    result_path = args.output_dir / "06_candidate_selection_results.jsonl"
    cache = read_result_cache(result_path)
    successful = sum(
        1
        for job in api_jobs
        if cache.get(job["job_id"], {}).get("status") == "success"
    )
    plan_path = write_plan(
        args.output_dir,
        retrieval,
        api_jobs,
        config,
        args.model,
        args.splits,
    )

    print("候補30件→10件選定の入力確認が完了しました。")
    print(f"  対象購入意図: {len(api_jobs)}件")
    print("  各購入意図の入力候補: 30件")
    print("  各購入意図の出力候補: 10件")
    print(f"  選定モデル: {args.model}")
    print(f"  既存の成功結果: {successful}/{len(api_jobs)}件")
    print(f"  対象商品seed: {args.seed}")
    print(f"  計画ファイル: {plan_path}")

    if args.execute:
        if not str(__import__("os").environ.get("OPENAI_API_KEY", "")).strip():
            raise EnvironmentError(
                "OPENAI_API_KEYがありません。API利用が承認されるまで設定・実行しないでください。"
            )
        execute_jobs(
            api_jobs,
            cache,
            result_path,
            args.max_retries,
        )
        cache = read_result_cache(result_path)
    else:
        print("API呼び出し: 0回")
        print(
            "計画のみ作成しました。API利用が明示的に承認されるまで "
            "--execute は付けないでください。"
        )

    all_success = all(
        cache.get(job["job_id"], {}).get("status") == "success"
        for job in api_jobs
    )
    if all_success and (args.execute or args.rebuild_final):
        assignments, instances = build_final_outputs(
            retrieval,
            intents,
            api_jobs,
            cache,
            args.seed,
        )
        output_paths = write_final_outputs(
            args.output_dir,
            assignments,
            instances,
            cache,
            args.seed,
            args.model,
        )
        print("候補10件と対象商品を確定しました。")
        for path in output_paths:
            print(f"  {path}")
    elif args.rebuild_final:
        raise ValueError(
            "成功結果がそろっていないため、最終候補ファイルを再生成できません。"
        )
    else:
        print(
            "候補10件と対象商品はまだ未確定です。"
            "全ジョブ成功後に自動生成されます。"
        )


if __name__ == "__main__":
    main()
