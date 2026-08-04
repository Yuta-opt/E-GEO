from __future__ import annotations

"""先行研究準拠のTOEIC版E-GEO結果分析。

- Testの初期版／最適化版を対応比較する。
- text-embedding-3-largeで15プロンプトの全評価版を埋め込み、
  重心距離とペアワイズ距離を追跡する。
- 10特徴は自動キーワード一致にせず、論文と同じ0/1/2の人手評価表を作る。

埋め込みAPIは --execute-embeddings を付けたときだけ呼び出す。
"""

import argparse
import hashlib
import importlib.util
import json
import os
import sys
from itertools import combinations
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


LEGACY_ANALYSIS_PATH = Path(__file__).with_name(
    "06_TOEICメタ最適化実験結果を分析.py"
)
DEFAULT_RUN_DIR = Path(
    "日本語版データ/TOEIC/05_API実験/02_OpenAI_Gemini_Claude実行"
)
DEFAULT_OUTPUT_DIR = Path(
    "日本語版データ/TOEIC/06_分析結果/02_OpenAI_Gemini_Claude"
)
DEFAULT_EMBEDDING_MODEL = "text-embedding-3-large"

FEATURES: list[tuple[str, str]] = [
    ("search_engine_goal", "検索エンジンで順位を上げる目標"),
    ("user_intent", "実際の利用者の検索意図との整合"),
    ("keywords_synonyms", "関連キーワード・同義語・関連語"),
    ("opening_summary", "商品・目的・主な利点の冒頭要約"),
    ("section_headings", "特徴や用途などの見出し構成"),
    ("scannable_bullets", "走査しやすい箇条書き・短い段落"),
    ("use_cases", "具体的な利用場面・適用例"),
    ("faq", "利用者が尋ねそうな質問と回答"),
    ("no_keyword_stuffing", "不自然なキーワード反復を避ける指示"),
    ("maintains_factuality", "元の事実を維持し未裏付け情報を加えない指示"),
]


def load_module(path: Path, module_name: str) -> Any:
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"モジュールを読み込めません: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


legacy = load_module(LEGACY_ANALYSIS_PATH, "toeic_egeo_legacy_analysis")


def read_json(path: Path) -> Any:
    if not path.exists():
        raise FileNotFoundError(f"ファイルが見つかりません: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"JSONLが見つかりません: {path}")
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(
        path.read_text(encoding="utf-8").splitlines(),
        start=1,
    ):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSONLを読めません: {path}:{line_number}") from exc
        if not isinstance(value, dict):
            raise ValueError(f"JSONL各行はobjectである必要があります: {line_number}")
        rows.append(value)
    return rows


def stable_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def prompt_trajectory(versions: list[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for item in versions:
        rows.append(
            {
                "prompt_id": str(item["prompt_id"]),
                "version_number": int(item["version_number"]),
                "trajectory_step": int(item["version_number"]) - 1,
                "epoch": int(item["epoch"]),
                "batch": int(item["batch"]),
                "version_label": str(item["version_label"]),
                "prompt_text": str(item["prompt_text"]),
                "validation_mean": float(
                    item["validation_summary"]["overall_mean"]
                ),
            }
        )
    frame = pd.DataFrame(rows)
    if frame.empty:
        raise ValueError("最適化版履歴が空です。")
    counts = frame.groupby("version_number")["prompt_id"].nunique()
    if not counts.eq(15).all():
        raise ValueError(
            "各評価版に15プロンプトがそろっていません: "
            + str(counts.to_dict())
        )
    return frame.sort_values(
        ["version_number", "prompt_id"],
        kind="stable",
    ).reset_index(drop=True)


def load_embedding_cache(path: Path) -> dict[str, list[float]]:
    cache: dict[str, list[float]] = {}
    if not path.exists():
        return cache
    for row in read_jsonl(path):
        cache[str(row["cache_key"])] = [
            float(value) for value in row["embedding"]
        ]
    return cache


def append_jsonl(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(row, ensure_ascii=False) + "\n")


def embedding_for_text(
    *,
    text: str,
    model: str,
    cache: dict[str, list[float]],
    cache_path: Path,
    execute: bool,
    client: Any | None,
) -> np.ndarray:
    key = stable_hash(model + "\n" + text)
    if key in cache:
        return np.asarray(cache[key], dtype=np.float64)
    if not execute:
        raise RuntimeError(
            "埋め込みキャッシュが不足しています。"
            "先行研究準拠の収束分析を実行するには "
            "--execute-embeddings を付けてください。"
        )
    if client is None:
        raise RuntimeError("OpenAI clientが初期化されていません。")
    response = client.embeddings.create(model=model, input=text)
    embedding = [float(value) for value in response.data[0].embedding]
    row = {
        "cache_key": key,
        "model": model,
        "text_sha256": stable_hash(text),
        "embedding": embedding,
        "usage": {
            "prompt_tokens": int(
                getattr(getattr(response, "usage", None), "prompt_tokens", 0)
                or 0
            ),
            "total_tokens": int(
                getattr(getattr(response, "usage", None), "total_tokens", 0)
                or 0
            ),
        },
    }
    append_jsonl(cache_path, row)
    cache[key] = embedding
    return np.asarray(embedding, dtype=np.float64)


def normalize_rows(matrix: np.ndarray) -> np.ndarray:
    norms = np.linalg.norm(matrix, axis=1, keepdims=True)
    return matrix / np.clip(norms, 1e-12, None)


def convergence_metrics(matrix: np.ndarray) -> dict[str, float]:
    normalized = normalize_rows(matrix)
    centroid = normalized.mean(axis=0)
    centroid = centroid / max(float(np.linalg.norm(centroid)), 1e-12)
    centroid_distances = 1.0 - normalized @ centroid
    pairwise = [
        1.0 - float(normalized[i] @ normalized[j])
        for i, j in combinations(range(len(normalized)), 2)
    ]
    return {
        "mean_cosine_distance_to_centroid": float(
            np.mean(centroid_distances)
        ),
        "mean_pairwise_cosine_distance": float(np.mean(pairwise)),
    }


def run_convergence_analysis(
    trajectory: pd.DataFrame,
    model: str,
    cache_path: Path,
    execute: bool,
) -> pd.DataFrame:
    load_dotenv()
    client = None
    if execute:
        if not os.environ.get("OPENAI_API_KEY", "").strip():
            raise EnvironmentError(
                "--execute-embeddingsにはOPENAI_API_KEYが必要です。"
            )
        from openai import OpenAI

        client = OpenAI()

    cache = load_embedding_cache(cache_path)
    rows: list[dict[str, Any]] = []
    for version_number, group in trajectory.groupby(
        "version_number",
        sort=True,
    ):
        group = group.sort_values("prompt_id", kind="stable")
        vectors = np.vstack(
            [
                embedding_for_text(
                    text=str(text),
                    model=model,
                    cache=cache,
                    cache_path=cache_path,
                    execute=execute,
                    client=client,
                )
                for text in group["prompt_text"]
            ]
        )
        metrics = convergence_metrics(vectors)
        first = group.iloc[0]
        rows.append(
            {
                "version_number": int(version_number),
                "trajectory_step": int(version_number) - 1,
                "epoch": int(first["epoch"]),
                "batch": int(first["batch"]),
                "prompt_count": int(len(group)),
                "embedding_model": model,
                **metrics,
            }
        )
    return pd.DataFrame(rows)


def create_feature_review_workbook(
    path: Path,
    final_prompts: list[dict[str, Any]],
) -> None:
    if path.exists():
        return

    score_rows: list[dict[str, Any]] = []
    for item in final_prompts:
        for condition, field in [
            ("initial", "initial_prompt"),
            ("optimized", "optimized_prompt"),
        ]:
            row: dict[str, Any] = {
                "prompt_id": str(item["prompt_id"]),
                "condition": condition,
            }
            for feature_id, _ in FEATURES:
                row[feature_id] = ""
            row["reviewer_notes"] = ""
            row["prompt_text"] = str(item[field])
            score_rows.append(row)

    instructions = pd.DataFrame(
        [
            ["評価単位", "15初期プロンプトと15最適化プロンプトの計30件"],
            ["0", "Absent：特徴が存在しない"],
            ["1", "Implicit：直接は書かれていないが、実質的に含まれる"],
            ["2", "Explicit：特徴が明示的に指示されている"],
            [
                "注意",
                "単語の有無だけで決めず、プロンプト全体の意味を読んで人手判定する。",
            ],
            [
                "論文対応",
                "E-GEO v2 Section 5.5.1の三段階人手評価に対応する。",
            ],
        ],
        columns=["項目", "説明"],
    )
    feature_definitions = pd.DataFrame(
        [
            [feature_id, label]
            for feature_id, label in FEATURES
        ],
        columns=["feature_id", "日本語定義"],
    )
    scores = pd.DataFrame(score_rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        instructions.to_excel(writer, sheet_name="instructions", index=False)
        feature_definitions.to_excel(
            writer,
            sheet_name="feature_definitions",
            index=False,
        )
        scores.to_excel(writer, sheet_name="scores", index=False)

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    score_sheet = workbook["scores"]
    score_sheet.column_dimensions["A"].width = 18
    score_sheet.column_dimensions["B"].width = 14
    for index in range(3, 3 + len(FEATURES)):
        score_sheet.column_dimensions[
            score_sheet.cell(row=1, column=index).column_letter
        ].width = 20
    score_sheet.column_dimensions[
        score_sheet.cell(row=1, column=3 + len(FEATURES)).column_letter
    ].width = 35
    score_sheet.column_dimensions[
        score_sheet.cell(row=1, column=4 + len(FEATURES)).column_letter
    ].width = 120
    workbook.save(path)


def load_manual_feature_scores(
    path: Path,
) -> tuple[pd.DataFrame, bool]:
    frame = pd.read_excel(path, sheet_name="scores", dtype=object).fillna("")
    feature_ids = [feature_id for feature_id, _ in FEATURES]
    complete = True
    for feature_id in feature_ids:
        values = pd.to_numeric(frame[feature_id], errors="coerce")
        valid = values.isin([0, 1, 2])
        if not valid.all():
            complete = False
        frame[feature_id] = values
    return frame, complete


def create_feature_outputs(
    scores: pd.DataFrame,
    output_dir: Path,
) -> tuple[Path, Path]:
    feature_ids = [feature_id for feature_id, _ in FEATURES]
    summary_rows: list[dict[str, Any]] = []
    for condition, group in scores.groupby("condition", sort=False):
        for feature_id, label in FEATURES:
            values = pd.to_numeric(group[feature_id], errors="raise")
            summary_rows.append(
                {
                    "condition": condition,
                    "feature_id": feature_id,
                    "feature_label_ja": label,
                    "prompt_count": int(len(group)),
                    "mean_score": float(values.mean()),
                    "absent_count": int((values == 0).sum()),
                    "implicit_count": int((values == 1).sum()),
                    "explicit_count": int((values == 2).sum()),
                }
            )
    summary_path = output_dir / "05_prompt_feature_summary.csv"
    pd.DataFrame(summary_rows).to_csv(
        summary_path,
        index=False,
        encoding="utf-8-sig",
    )

    ordered = scores.copy()
    ordered["row_label"] = ordered["prompt_id"] + " / " + ordered["condition"]
    matrix = ordered.set_index("row_label")[feature_ids].astype(float)
    plt.figure(figsize=(14, 12))
    plt.imshow(matrix.to_numpy(), aspect="auto", vmin=0, vmax=2)
    plt.colorbar(label="0=Absent, 1=Implicit, 2=Explicit")
    plt.xticks(
        range(len(feature_ids)),
        [label for _, label in FEATURES],
        rotation=45,
        ha="right",
    )
    plt.yticks(range(len(matrix)), matrix.index)
    plt.title("15初期・15最適化プロンプトの10特徴人手評価")
    plt.tight_layout()
    heatmap_path = output_dir / "11_manual_feature_heatmap.png"
    plt.savefig(heatmap_path, dpi=200)
    plt.close()
    return summary_path, heatmap_path


def create_test_charts(
    paired: pd.DataFrame,
    test_summary: pd.DataFrame,
    output_dir: Path,
) -> list[Path]:
    paths: list[Path] = []
    if not paired.empty:
        chart = paired.groupby("prompt_id")[["initial_mean", "optimized_mean"]].mean()
        axis = chart.plot(kind="bar", figsize=(13, 6))
        axis.set_title("Test長文：初期プロンプトと最適化プロンプト")
        axis.set_xlabel("初期プロンプト")
        axis.set_ylabel("平均順位改善量")
        axis.axhline(0, linewidth=1)
        plt.tight_layout()
        path = output_dir / "07_initial_vs_optimized.png"
        plt.savefig(path, dpi=200)
        plt.close()
        paths.append(path)

    optimized = test_summary.loc[
        test_summary["condition"].isin(
            ["test_optimized_long", "test_optimized_short"]
        )
    ]
    if not optimized.empty:
        chart = optimized.groupby(["prompt_id", "query_form"])[
            "mean_rank_improvement"
        ].mean().unstack("query_form")
        axis = chart.plot(kind="bar", figsize=(13, 6))
        axis.set_title("最適化プロンプト：長文Testと短文Test")
        axis.set_xlabel("初期プロンプト")
        axis.set_ylabel("平均順位改善量")
        axis.axhline(0, linewidth=1)
        plt.tight_layout()
        path = output_dir / "08_long_vs_short.png"
        plt.savefig(path, dpi=200)
        plt.close()
        paths.append(path)
    return paths


def create_convergence_chart(
    convergence: pd.DataFrame,
    output_dir: Path,
) -> Path:
    axis = convergence.plot(
        x="trajectory_step",
        y=[
            "mean_cosine_distance_to_centroid",
            "mean_pairwise_cosine_distance",
        ],
        marker="o",
        figsize=(9, 5),
    )
    axis.set_title("15プロンプトの埋め込み空間における収束")
    axis.set_xlabel("評価ステップ（0=初期版）")
    axis.set_ylabel("コサイン距離")
    axis.grid(True, alpha=0.3)
    plt.tight_layout()
    path = output_dir / "09_prompt_convergence.png"
    plt.savefig(path, dpi=200)
    plt.close()
    return path


def write_review_workbook(
    path: Path,
    test_summary: pd.DataFrame,
    paired: pd.DataFrame,
    convergence: pd.DataFrame,
    trajectory: pd.DataFrame,
    final_prompts: list[dict[str, Any]],
    manual_feature_complete: bool,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        test_summary.to_excel(writer, sheet_name="test_summary", index=False)
        paired.to_excel(writer, sheet_name="paired_test", index=False)
        convergence.to_excel(writer, sheet_name="convergence", index=False)
        trajectory.to_excel(writer, sheet_name="prompt_trajectory", index=False)
        pd.DataFrame(final_prompts).to_excel(
            writer,
            sheet_name="final_prompts",
            index=False,
        )
        pd.DataFrame(
            [
                {
                    "manual_feature_scoring_complete": manual_feature_complete,
                    "meaning": (
                        "Falseの場合は04_prompt_feature_manual_review.xlsxのscoresを"
                        "0/1/2で人手入力して再実行する。"
                    ),
                }
            ]
        ).to_excel(writer, sheet_name="analysis_status", index=False)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="先行研究準拠の統計・埋め込み収束・人手特徴評価表を作る。"
    )
    parser.add_argument("--run-dir", type=Path, default=DEFAULT_RUN_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--embedding-model", default=DEFAULT_EMBEDDING_MODEL)
    parser.add_argument("--execute-embeddings", action="store_true")
    args = parser.parse_args()

    versions = read_jsonl(args.run_dir / "02_optimization_versions.jsonl")
    final_prompts = read_json(args.run_dir / "03_final_prompts.json")
    test_rows = read_jsonl(args.run_dir / "04_test_results.jsonl")
    run_summary = read_json(args.run_dir / "06_run_summary.json")

    if not isinstance(final_prompts, list) or len(final_prompts) != 15:
        raise ValueError("正式分析には15件のfinal promptsが必要です。")
    if not test_rows:
        raise ValueError("Test結果が空です。")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    test_frame = pd.DataFrame(test_rows)
    test_frame["rank_improvement"] = pd.to_numeric(
        test_frame["rank_improvement"],
        errors="raise",
    )
    test_summary = legacy.summarize_test(test_frame)
    paired = legacy.paired_initial_optimized(test_frame)
    trajectory = prompt_trajectory(versions)
    convergence = run_convergence_analysis(
        trajectory,
        args.embedding_model,
        args.output_dir / "11_embedding_cache.jsonl",
        args.execute_embeddings,
    )

    test_summary.to_csv(
        args.output_dir / "01_test_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    paired.to_csv(
        args.output_dir / "02_initial_vs_optimized_paired.csv",
        index=False,
        encoding="utf-8-sig",
    )
    convergence.to_csv(
        args.output_dir / "03_prompt_convergence_trajectory.csv",
        index=False,
        encoding="utf-8-sig",
    )
    trajectory.to_csv(
        args.output_dir / "03b_prompt_trajectory.csv",
        index=False,
        encoding="utf-8-sig",
    )

    feature_review_path = args.output_dir / "04_prompt_feature_manual_review.xlsx"
    create_feature_review_workbook(feature_review_path, final_prompts)
    feature_scores, manual_feature_complete = load_manual_feature_scores(
        feature_review_path
    )
    feature_outputs: list[str] = []
    if manual_feature_complete:
        feature_outputs = [
            str(path)
            for path in create_feature_outputs(feature_scores, args.output_dir)
        ]

    chart_paths = create_test_charts(test_summary, paired, args.output_dir)
    chart_paths.append(create_convergence_chart(convergence, args.output_dir))

    review_path = args.output_dir / "10_results_review.xlsx"
    write_review_workbook(
        review_path,
        test_summary,
        paired,
        convergence,
        trajectory,
        final_prompts,
        manual_feature_complete,
    )

    summary = {
        "analysis_status": (
            "complete"
            if manual_feature_complete
            else "numeric_and_embedding_complete_manual_feature_scoring_pending"
        ),
        "run_status": run_summary.get("status"),
        "embedding_model": args.embedding_model,
        "embedding_trajectory_steps": int(len(convergence)),
        "manual_feature_scoring_complete": manual_feature_complete,
        "manual_feature_scale": {
            "0": "absent",
            "1": "implicit",
            "2": "explicit",
        },
        "automatic_keyword_feature_scoring_used": false,
        "charts": [str(path) for path in chart_paths],
        "feature_outputs": feature_outputs,
        "prior_study_alignment": {
            "embedding_model_matches_paper": (
                args.embedding_model == "text-embedding-3-large"
            ),
            "feature_scoring_method": "manual three-level review",
        },
    }
    (args.output_dir / "06_results_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("06b 先行研究準拠分析完了")
    print(f"  Test集計: {len(test_summary)}行")
    print(f"  対応比較: {len(paired)}行")
    print(f"  収束ステップ: {len(convergence)}")
    print(f"  埋め込みモデル: {args.embedding_model}")
    print(
        "  10特徴人手評価: "
        + ("完了" if manual_feature_complete else "未入力（評価表を作成済み）")
    )
    print(f"  出力先: {args.output_dir}")


if __name__ == "__main__":
    main()
