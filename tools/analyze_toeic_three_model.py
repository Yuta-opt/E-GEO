from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.stats import rankdata, wilcoxon
from statsmodels.stats.multitest import multipletests


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MAIN_RESULTS = (
    REPO_ROOT
    / "日本語版データ"
    / "TOEIC"
    / "05_API実験"
    / "02_OpenAI_Gemini_Claude実行"
    / "04_test_results.jsonl"
)
DEFAULT_CLAUDE_RESULTS = (
    REPO_ROOT
    / "日本語版データ"
    / "TOEIC"
    / "05_API実験"
    / "03_Claude_Heldout評価"
    / "02_claude_test_results.jsonl"
)
DEFAULT_OUTPUT_DIR = (
    REPO_ROOT
    / "日本語版データ"
    / "TOEIC"
    / "06_分析結果"
    / "03_3モデル統合"
)

MODEL_DISPLAY = {
    "Model E": "GPT-5",
    "Model F": "Gemini 3.5 Flash",
    "Model G": "Claude Sonnet 4.5",
}
MODEL_ORDER = ["Model E", "Model F", "Model G"]
CONDITIONS = ["test_initial_long", "test_optimized_long"]
BOOTSTRAP_RESAMPLES = 10_000


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"JSONLが見つかりません: {path}")
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSONLを読めません: {path}:{line_number}") from exc
        if not isinstance(row, dict):
            raise ValueError(f"JSONL各行はobjectである必要があります: {path}:{line_number}")
        rows.append(row)
    return rows


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def stable_seed(*parts: str) -> int:
    raw = "\n".join(parts).encode("utf-8")
    return int(hashlib.sha256(raw).hexdigest()[:8], 16)


def bootstrap_mean_ci(values: np.ndarray, *, seed: int) -> tuple[float, float]:
    clean = np.asarray(values, dtype=np.float64)
    clean = clean[np.isfinite(clean)]
    if clean.size == 0:
        return float("nan"), float("nan")
    if clean.size == 1:
        value = float(clean[0])
        return value, value
    rng = np.random.default_rng(seed)
    indexes = rng.integers(0, clean.size, size=(BOOTSTRAP_RESAMPLES, clean.size))
    means = clean[indexes].mean(axis=1)
    low, high = np.quantile(means, [0.025, 0.975])
    return float(low), float(high)


def sem(values: np.ndarray) -> float:
    clean = np.asarray(values, dtype=np.float64)
    clean = clean[np.isfinite(clean)]
    if clean.size <= 1:
        return 0.0
    return float(clean.std(ddof=1) / math.sqrt(clean.size))


def safe_wilcoxon(values: np.ndarray) -> tuple[float, float, int]:
    clean = np.asarray(values, dtype=np.float64)
    clean = clean[np.isfinite(clean)]
    nonzero = clean[~np.isclose(clean, 0.0)]
    if nonzero.size == 0:
        return 0.0, 1.0, 0
    statistic, p_value = wilcoxon(
        nonzero,
        zero_method="wilcox",
        alternative="two-sided",
    )
    return float(statistic), float(p_value), int(nonzero.size)


def matched_rank_biserial(values: np.ndarray) -> float:
    clean = np.asarray(values, dtype=np.float64)
    clean = clean[np.isfinite(clean)]
    nonzero = clean[~np.isclose(clean, 0.0)]
    if nonzero.size == 0:
        return 0.0
    ranks = rankdata(np.abs(nonzero), method="average")
    positive = float(ranks[nonzero > 0].sum())
    negative = float(ranks[nonzero < 0].sum())
    denominator = positive + negative
    return 0.0 if denominator == 0 else (positive - negative) / denominator


def load_and_validate(
    main_path: Path,
    claude_path: Path,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    main_rows = read_jsonl(main_path)
    claude_rows = read_jsonl(claude_path)
    main = pd.DataFrame(main_rows)
    claude = pd.DataFrame(claude_rows)

    required = {
        "condition",
        "query_form",
        "prompt_id",
        "intent_id",
        "reranker_label",
        "reranker_model",
        "rank_improvement",
    }
    for label, frame in [("OpenAI/Gemini", main), ("Claude", claude)]:
        missing_columns = sorted(required - set(frame.columns))
        if missing_columns:
            raise ValueError(f"{label}結果の列が不足しています: {missing_columns}")

    main_long = main.loc[
        main["condition"].isin(CONDITIONS)
        & main["reranker_label"].isin(MODEL_ORDER[:2])
    ].copy()
    claude_long = claude.loc[
        claude["condition"].isin(CONDITIONS)
        & (claude["reranker_label"] == "Model G")
    ].copy()
    frame = pd.concat([main_long, claude_long], ignore_index=True)
    frame["rank_improvement"] = pd.to_numeric(
        frame["rank_improvement"],
        errors="raise",
    )
    frame["prompt_id"] = frame["prompt_id"].astype(str)
    frame["intent_id"] = frame["intent_id"].astype(str)
    frame["reranker_label"] = frame["reranker_label"].astype(str)
    frame["model_display"] = frame["reranker_label"].map(MODEL_DISPLAY)
    if frame["model_display"].isna().any():
        unknown = sorted(
            frame.loc[frame["model_display"].isna(), "reranker_label"].unique()
        )
        raise ValueError(f"未知のモデルラベルです: {unknown}")

    duplicate_columns = [
        "condition",
        "prompt_id",
        "intent_id",
        "reranker_label",
    ]
    duplicates = frame.duplicated(duplicate_columns, keep=False)
    if duplicates.any():
        sample = frame.loc[duplicates, duplicate_columns].head(10).to_dict("records")
        raise ValueError(f"統合結果に重複があります: {sample}")

    counts: dict[str, Any] = {}
    for model_label in MODEL_ORDER:
        model = frame.loc[frame["reranker_label"] == model_label]
        condition_counts = model["condition"].value_counts().to_dict()
        prompt_count = int(model["prompt_id"].nunique())
        intent_count = int(model["intent_id"].nunique())
        counts[model_label] = {
            "display": MODEL_DISPLAY[model_label],
            "rows": int(len(model)),
            "condition_counts": {
                str(key): int(value) for key, value in condition_counts.items()
            },
            "prompt_count": prompt_count,
            "intent_count": intent_count,
            "model_ids": sorted(
                str(value) for value in model["reranker_model"].dropna().unique()
            ),
        }
        expected = {
            "test_initial_long": 450,
            "test_optimized_long": 450,
        }
        if (
            condition_counts != expected
            or prompt_count != 15
            or intent_count != 30
        ):
            raise ValueError(
                f"{model_label}の件数が不正です: counts={condition_counts}, "
                f"prompts={prompt_count}, intents={intent_count}"
            )

    if len(frame) != 2700:
        raise ValueError(f"3モデル統合行数が2700ではありません: {len(frame)}")

    validation = {
        "status": "pass",
        "main_source_rows": int(len(main)),
        "claude_source_rows": int(len(claude)),
        "integrated_long_rows": int(len(frame)),
        "models": counts,
        "api_calls": 0,
    }
    return frame, validation


def build_paired(frame: pd.DataFrame) -> pd.DataFrame:
    pivot = frame.pivot(
        index=[
            "reranker_label",
            "model_display",
            "reranker_model",
            "prompt_id",
            "intent_id",
        ],
        columns="condition",
        values="rank_improvement",
    ).reset_index()
    if pivot[CONDITIONS].isna().any().any():
        raise ValueError("初期・最適化の対応ペアに欠損があります。")
    pivot["recovery_optimized_minus_initial"] = (
        pivot["test_optimized_long"] - pivot["test_initial_long"]
    )
    if len(pivot) != 1350:
        raise ValueError(f"対応ペアが1350行ではありません: {len(pivot)}")
    return pivot


def summarize_one_unit(
    unit: pd.DataFrame,
    *,
    label: str,
    display: str,
    model_id: str,
    model_count: int,
) -> dict[str, Any]:
    initial = unit["initial_mean"].to_numpy(dtype=np.float64)
    optimized = unit["optimized_mean"].to_numpy(dtype=np.float64)
    recovery = unit["recovery_mean"].to_numpy(dtype=np.float64)
    initial_low, initial_high = bootstrap_mean_ci(
        initial,
        seed=stable_seed(label, "initial"),
    )
    optimized_low, optimized_high = bootstrap_mean_ci(
        optimized,
        seed=stable_seed(label, "optimized"),
    )
    recovery_low, recovery_high = bootstrap_mean_ci(
        recovery,
        seed=stable_seed(label, "recovery"),
    )
    statistic, p_value, nonzero_n = safe_wilcoxon(recovery)
    return {
        "model_label": label,
        "model_display": display,
        "model_id": model_id,
        "model_count": model_count,
        "n_prompts": int(len(unit)),
        "n_intents_per_prompt": 30,
        "n_prompt_intent_pairs": int(len(unit) * 30 * model_count),
        "initial_mean": float(initial.mean()),
        "initial_sem_prompt_clustered": sem(initial),
        "initial_ci95_low_prompt_bootstrap": initial_low,
        "initial_ci95_high_prompt_bootstrap": initial_high,
        "optimized_mean": float(optimized.mean()),
        "optimized_sem_prompt_clustered": sem(optimized),
        "optimized_ci95_low_prompt_bootstrap": optimized_low,
        "optimized_ci95_high_prompt_bootstrap": optimized_high,
        "recovery_mean": float(recovery.mean()),
        "recovery_sem_prompt_clustered": sem(recovery),
        "recovery_ci95_low_prompt_bootstrap": recovery_low,
        "recovery_ci95_high_prompt_bootstrap": recovery_high,
        "prompts_improved": int((recovery > 0).sum()),
        "prompts_unchanged": int(np.isclose(recovery, 0.0).sum()),
        "prompts_regressed": int((recovery < 0).sum()),
        "wilcoxon_unit": "prompt mean across 30 Test intents",
        "wilcoxon_n": int(len(recovery)),
        "wilcoxon_nonzero_n": nonzero_n,
        "wilcoxon_statistic": statistic,
        "wilcoxon_p_value": p_value,
        "matched_rank_biserial": matched_rank_biserial(recovery),
        "ci_method": (
            f"percentile bootstrap by prompt, "
            f"{BOOTSTRAP_RESAMPLES} resamples"
        ),
    }


def model_level_summary(
    paired: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    prompt_means = (
        paired.groupby(
            [
                "reranker_label",
                "model_display",
                "reranker_model",
                "prompt_id",
            ],
            as_index=False,
        )
        .agg(
            initial_mean=("test_initial_long", "mean"),
            optimized_mean=("test_optimized_long", "mean"),
            recovery_mean=("recovery_optimized_minus_initial", "mean"),
            recovery_median=("recovery_optimized_minus_initial", "median"),
            intent_pair_improved_rate=(
                "recovery_optimized_minus_initial",
                lambda series: float((series > 0).mean()),
            ),
        )
    )

    rows: list[dict[str, Any]] = []
    for model_label in MODEL_ORDER:
        group = prompt_means.loc[
            prompt_means["reranker_label"] == model_label
        ].copy()
        model_ids = sorted(group["reranker_model"].astype(str).unique())
        rows.append(
            summarize_one_unit(
                group,
                label=model_label,
                display=MODEL_DISPLAY[model_label],
                model_id=" | ".join(model_ids),
                model_count=1,
            )
        )

    overall_prompt = (
        prompt_means.groupby("prompt_id", as_index=False)
        .agg(
            initial_mean=("initial_mean", "mean"),
            optimized_mean=("optimized_mean", "mean"),
            recovery_mean=("recovery_mean", "mean"),
        )
    )
    rows.append(
        summarize_one_unit(
            overall_prompt,
            label="Overall",
            display="3モデル平均",
            model_id="GPT-5 | Gemini 3.5 Flash | Claude Sonnet 4.5",
            model_count=3,
        )
    )
    summary = pd.DataFrame(rows)

    model_mask = summary["model_label"].isin(MODEL_ORDER)
    model_p = summary.loc[model_mask, "wilcoxon_p_value"].to_numpy(dtype=float)
    _, adjusted, _, _ = multipletests(
        model_p,
        alpha=0.05,
        method="holm",
    )
    summary["wilcoxon_p_holm_three_models"] = np.nan
    summary["significant_holm_0_05"] = False
    summary.loc[
        model_mask,
        "wilcoxon_p_holm_three_models",
    ] = adjusted
    summary.loc[
        model_mask,
        "significant_holm_0_05",
    ] = adjusted < 0.05
    overall_mask = summary["model_label"] == "Overall"
    summary.loc[overall_mask, "significant_holm_0_05"] = (
        summary.loc[overall_mask, "wilcoxon_p_value"] < 0.05
    )
    summary["multiple_testing_note"] = np.where(
        overall_mask,
        "Overall is the prespecified aggregate test; no Holm adjustment applied.",
        "Holm correction across the three model-specific tests.",
    )

    order = {
        label: index
        for index, label in enumerate(MODEL_ORDER + ["Overall"])
    }
    summary["_order"] = summary["model_label"].map(order)
    summary = (
        summary.sort_values("_order")
        .drop(columns="_order")
        .reset_index(drop=True)
    )
    return summary, prompt_means


def prompt_level_tests(paired: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    grouped = paired.groupby(
        ["reranker_label", "model_display", "prompt_id"],
        sort=False,
    )
    for (model_label, display, prompt_id), group in grouped:
        values = group["recovery_optimized_minus_initial"].to_numpy(
            dtype=np.float64
        )
        low, high = bootstrap_mean_ci(
            values,
            seed=stable_seed(
                str(model_label),
                str(prompt_id),
                "intent_bootstrap",
            ),
        )
        statistic, p_value, nonzero_n = safe_wilcoxon(values)
        rows.append(
            {
                "model_label": model_label,
                "model_display": display,
                "prompt_id": prompt_id,
                "n_intents": int(len(values)),
                "initial_mean": float(group["test_initial_long"].mean()),
                "optimized_mean": float(group["test_optimized_long"].mean()),
                "recovery_mean": float(values.mean()),
                "recovery_median": float(np.median(values)),
                "recovery_ci95_low_intent_bootstrap": low,
                "recovery_ci95_high_intent_bootstrap": high,
                "improved_intent_rate": float((values > 0).mean()),
                "wilcoxon_nonzero_n": nonzero_n,
                "wilcoxon_statistic": statistic,
                "wilcoxon_p_value": p_value,
                "matched_rank_biserial": matched_rank_biserial(values),
            }
        )
    output = pd.DataFrame(rows)
    all_p = output["wilcoxon_p_value"].to_numpy(dtype=float)
    _, global_adjusted, _, _ = multipletests(
        all_p,
        alpha=0.05,
        method="holm",
    )
    output["wilcoxon_p_holm_global_45"] = global_adjusted
    output["significant_holm_global_0_05"] = global_adjusted < 0.05
    output["wilcoxon_p_holm_within_model_15"] = np.nan
    output["significant_holm_within_model_0_05"] = False
    for _, indexes in output.groupby("model_label").groups.items():
        index_list = list(indexes)
        _, adjusted, _, _ = multipletests(
            output.loc[index_list, "wilcoxon_p_value"].to_numpy(dtype=float),
            alpha=0.05,
            method="holm",
        )
        output.loc[
            index_list,
            "wilcoxon_p_holm_within_model_15",
        ] = adjusted
        output.loc[
            index_list,
            "significant_holm_within_model_0_05",
        ] = adjusted < 0.05
    order = {
        label: index for index, label in enumerate(MODEL_ORDER)
    }
    output["_order"] = output["model_label"].map(order)
    return (
        output.sort_values(["_order", "prompt_id"])
        .drop(columns="_order")
        .reset_index(drop=True)
    )


def configure_japanese_font() -> None:
    plt.rcParams["font.sans-serif"] = [
        "Yu Gothic",
        "Meiryo",
        "Noto Sans CJK JP",
        "IPAexGothic",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def create_graphs(
    summary: pd.DataFrame,
    prompt_means: pd.DataFrame,
    output_dir: Path,
) -> list[Path]:
    configure_japanese_font()
    model_summary = summary.loc[
        summary["model_label"].isin(MODEL_ORDER)
    ].copy()
    labels = model_summary["model_display"].tolist()
    x = np.arange(len(labels))
    width = 0.35

    figure, axis = plt.subplots(figsize=(10, 6))
    initial_error = np.vstack(
        [
            model_summary["initial_mean"]
            - model_summary["initial_ci95_low_prompt_bootstrap"],
            model_summary["initial_ci95_high_prompt_bootstrap"]
            - model_summary["initial_mean"],
        ]
    )
    optimized_error = np.vstack(
        [
            model_summary["optimized_mean"]
            - model_summary["optimized_ci95_low_prompt_bootstrap"],
            model_summary["optimized_ci95_high_prompt_bootstrap"]
            - model_summary["optimized_mean"],
        ]
    )
    axis.bar(
        x - width / 2,
        model_summary["initial_mean"],
        width,
        yerr=initial_error,
        capsize=4,
        label="初期",
    )
    axis.bar(
        x + width / 2,
        model_summary["optimized_mean"],
        width,
        yerr=optimized_error,
        capsize=4,
        label="最適化後",
    )
    axis.axhline(0, linewidth=1)
    axis.set_xticks(x, labels)
    axis.set_ylabel("平均順位改善量（正の値ほど順位上昇）")
    axis.set_title("3評価モデルにおける初期・最適化プロンプトの比較")
    axis.legend()
    axis.grid(axis="y", alpha=0.25)
    figure.tight_layout()
    path1 = output_dir / "07_three_model_initial_vs_optimized.png"
    figure.savefig(path1, dpi=200, bbox_inches="tight")
    plt.close(figure)

    figure, axis = plt.subplots(figsize=(9, 5.5))
    recovery_error = np.vstack(
        [
            model_summary["recovery_mean"]
            - model_summary["recovery_ci95_low_prompt_bootstrap"],
            model_summary["recovery_ci95_high_prompt_bootstrap"]
            - model_summary["recovery_mean"],
        ]
    )
    axis.bar(
        labels,
        model_summary["recovery_mean"],
        yerr=recovery_error,
        capsize=5,
    )
    axis.axhline(0, linewidth=1)
    axis.set_ylabel("最適化による平均回復量")
    axis.set_title("最適化効果のモデル間比較")
    axis.grid(axis="y", alpha=0.25)
    figure.tight_layout()
    path2 = output_dir / "08_three_model_recovery.png"
    figure.savefig(path2, dpi=200, bbox_inches="tight")
    plt.close(figure)

    heat = prompt_means.pivot(
        index="prompt_id",
        columns="reranker_label",
        values="recovery_mean",
    )
    heat = heat.reindex(columns=MODEL_ORDER).sort_index()
    matrix = heat.to_numpy(dtype=float)
    limit = max(
        abs(float(np.nanmin(matrix))),
        abs(float(np.nanmax(matrix))),
        0.1,
    )
    figure, axis = plt.subplots(figsize=(8.5, 9))
    image = axis.imshow(
        matrix,
        aspect="auto",
        cmap="RdBu_r",
        vmin=-limit,
        vmax=limit,
    )
    axis.set_xticks(
        np.arange(len(MODEL_ORDER)),
        [MODEL_DISPLAY[label] for label in MODEL_ORDER],
    )
    axis.set_yticks(
        np.arange(len(heat.index)),
        heat.index.tolist(),
    )
    axis.set_xlabel("評価モデル")
    axis.set_ylabel("初期プロンプト")
    axis.set_title("プロンプト別の最適化回復量")
    for row_index in range(matrix.shape[0]):
        for column_index in range(matrix.shape[1]):
            value = matrix[row_index, column_index]
            axis.text(
                column_index,
                row_index,
                f"{value:.2f}",
                ha="center",
                va="center",
                fontsize=8,
            )
    colorbar = figure.colorbar(image, ax=axis)
    colorbar.set_label("最適化後 − 初期")
    figure.tight_layout()
    path3 = output_dir / "09_prompt_recovery_heatmap.png"
    figure.savefig(path3, dpi=200, bbox_inches="tight")
    plt.close(figure)
    return [path1, path2, path3]


def write_excel(
    path: Path,
    summary: pd.DataFrame,
    prompt_tests: pd.DataFrame,
    prompt_means: pd.DataFrame,
    paired: pd.DataFrame,
    validation: dict[str, Any],
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(
            writer,
            sheet_name="integrated_table",
            index=False,
        )
        summary.to_excel(
            writer,
            sheet_name="model_level_tests",
            index=False,
        )
        prompt_tests.to_excel(
            writer,
            sheet_name="prompt_level_tests",
            index=False,
        )
        prompt_means.to_excel(
            writer,
            sheet_name="prompt_by_model",
            index=False,
        )
        paired.to_excel(
            writer,
            sheet_name="paired_rows",
            index=False,
        )
        pd.DataFrame([validation]).to_excel(
            writer,
            sheet_name="validation",
            index=False,
        )

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        for column_cells in sheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column_cells) + 2,
                45,
            )
            sheet.column_dimensions[
                column_cells[0].column_letter
            ].width = width
    workbook.save(path)


def markdown_report(
    summary: pd.DataFrame,
    prompt_tests: pd.DataFrame,
) -> str:
    rows: list[str] = []
    for item in summary.itertuples(index=False):
        p_display = f"{item.wilcoxon_p_value:.4g}"
        adjusted = (
            "—"
            if pd.isna(item.wilcoxon_p_holm_three_models)
            else f"{item.wilcoxon_p_holm_three_models:.4g}"
        )
        rows.append(
            f"| {item.model_display} | {item.initial_mean:.4f} | "
            f"{item.optimized_mean:.4f} | {item.recovery_mean:+.4f} "
            f"[{item.recovery_ci95_low_prompt_bootstrap:+.4f}, "
            f"{item.recovery_ci95_high_prompt_bootstrap:+.4f}] | "
            f"{p_display} | {adjusted} |"
        )
    significant_global = int(
        prompt_tests["significant_holm_global_0_05"].sum()
    )
    significant_within = int(
        prompt_tests["significant_holm_within_model_0_05"].sum()
    )
    overall = summary.loc[
        summary["model_label"] == "Overall"
    ].iloc[0]
    return f"""# TOEIC E-GEO 3モデル統合分析

## 統合表

| 評価モデル | 初期平均 | 最適化後平均 | 回復量（95% CI） | Wilcoxon p | Holm補正p |
|---|---:|---:|---:|---:|---:|
{chr(10).join(rows)}

## 統計手法

- 主要な推論単位は**プロンプト**です。各プロンプトについて30件のTest intentを平均し、初期版と最適化版の差を15プロンプトで対応比較しました。
- モデル別検定はWilcoxon符号付順位検定を用い、3モデルのp値をHolm法で補正しました。
- 95%信頼区間はプロンプト単位のpercentile bootstrap（{BOOTSTRAP_RESAMPLES:,}回）です。
- 補助分析として、各モデル×各プロンプトの30 intent差にもWilcoxon検定を行い、全45比較およびモデル内15比較でHolm補正しました。

## 結果の読み方

3モデル平均では、初期プロンプトの平均順位改善量は{overall['initial_mean']:.4f}、最適化後は{overall['optimized_mean']:.4f}で、平均回復量は{overall['recovery_mean']:+.4f}でした。回復量の95% CIは[{overall['recovery_ci95_low_prompt_bootstrap']:+.4f}, {overall['recovery_ci95_high_prompt_bootstrap']:+.4f}]です。

補助的な45件のプロンプト別検定では、Holm補正後に有意だった比較は全体補正で{significant_global}件、モデル内補正で{significant_within}件でした。

**表現上の注意:** 最適化後の平均値は3モデルとも依然として0未満なので、「順位が全体として上昇した」ではなく、「初期プロンプトによる順位悪化が大幅に緩和された」と記述します。
"""


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "TOEIC E-GEOのGPT・Gemini・ClaudeをAPIなしで統合分析する。"
        )
    )
    parser.add_argument(
        "--mode",
        choices=["validate", "analyze"],
        default="validate",
    )
    parser.add_argument(
        "--main-results",
        type=Path,
        default=DEFAULT_MAIN_RESULTS,
    )
    parser.add_argument(
        "--claude-results",
        type=Path,
        default=DEFAULT_CLAUDE_RESULTS,
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
    )
    args = parser.parse_args()

    frame, validation = load_and_validate(
        args.main_results,
        args.claude_results,
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_json(
        args.output_dir / "00_validation_report.json",
        validation,
    )

    print("TOEIC E-GEO 3モデル統合分析")
    print(f"  Mode: {args.mode}")
    print(f"  Integrated rows: {len(frame)}")
    print("  Models: GPT-5 / Gemini 3.5 Flash / Claude Sonnet 4.5")
    print("  API calls: 0")

    if args.mode == "validate":
        print(
            "Validate完了：入力件数・対応関係を確認しました。"
            "分析ファイルは未生成です。"
        )
        return

    paired = build_paired(frame)
    summary, prompt_means = model_level_summary(paired)
    prompt_tests = prompt_level_tests(paired)

    summary.to_csv(
        args.output_dir / "01_three_model_integrated_table.csv",
        index=False,
        encoding="utf-8-sig",
    )
    summary.to_csv(
        args.output_dir / "02_model_level_statistical_tests.csv",
        index=False,
        encoding="utf-8-sig",
    )
    prompt_tests.to_csv(
        args.output_dir / "03_prompt_level_statistical_tests.csv",
        index=False,
        encoding="utf-8-sig",
    )
    prompt_means.to_csv(
        args.output_dir / "04_prompt_by_model_recovery.csv",
        index=False,
        encoding="utf-8-sig",
    )
    paired.to_csv(
        args.output_dir / "05_three_model_paired_rows.csv",
        index=False,
        encoding="utf-8-sig",
    )
    graph_paths = create_graphs(
        summary,
        prompt_means,
        args.output_dir,
    )
    write_excel(
        args.output_dir / "10_three_model_results.xlsx",
        summary,
        prompt_tests,
        prompt_means,
        paired,
        validation,
    )
    report = markdown_report(summary, prompt_tests)
    (args.output_dir / "11_three_model_report.md").write_text(
        report,
        encoding="utf-8",
    )

    overall = summary.loc[
        summary["model_label"] == "Overall"
    ].iloc[0]
    result_summary = {
        "status": "complete",
        "api_calls": 0,
        "integrated_source_rows": int(len(frame)),
        "paired_rows": int(len(paired)),
        "model_count": 3,
        "prompt_count": 15,
        "intent_count": 30,
        "primary_inference_unit": (
            "prompt mean across 30 Test intents"
        ),
        "overall_initial_mean": float(overall["initial_mean"]),
        "overall_optimized_mean": float(overall["optimized_mean"]),
        "overall_recovery_mean": float(overall["recovery_mean"]),
        "overall_recovery_ci95": [
            float(overall["recovery_ci95_low_prompt_bootstrap"]),
            float(overall["recovery_ci95_high_prompt_bootstrap"]),
        ],
        "overall_wilcoxon_p_value": float(
            overall["wilcoxon_p_value"]
        ),
        "model_specific_holm_significant_count": int(
            summary.loc[
                summary["model_label"].isin(MODEL_ORDER),
                "significant_holm_0_05",
            ].sum()
        ),
        "prompt_level_holm_global_significant_count": int(
            prompt_tests["significant_holm_global_0_05"].sum()
        ),
        "graphs": [str(path) for path in graph_paths],
        "poster_updated": False,
    }
    write_json(
        args.output_dir / "06_three_model_results_summary.json",
        result_summary,
    )

    print(json.dumps(result_summary, ensure_ascii=False, indent=2))
    print("3モデル統合表・統計検定・グラフ作成完了")
    print(f"  出力先: {args.output_dir}")
    print("  ポスター更新: なし")


if __name__ == "__main__":
    main()
